"""Tests for the 6.2 features (no audio device, no network)."""
import datetime
import io
import json
import os
import struct
import sys
import time
import zipfile
import zlib


def _tiny_pdf(path, text):
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R "
            "/Resources << /Font << /F1 5 0 R >> >> >>"]
    stream = f"BT /F1 14 Tf 20 100 Td ({text}) Tj ET"
    objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offs).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF".encode()
    with open(path, "wb") as f:
        f.write(out)


def _cfg(app, **kw):
    return app.sanitize(kw)


# ---- settings -----------------------------------------------------------------------------
def test_new_settings_are_checked(app):
    c = app.sanitize({})
    assert c["about_source"] == "text" and c["meeting_mode"] == "general" and c["search_days"] == "90"
    assert c["hide_from_share"] is False and c["diarize"] is False
    c = app.sanitize({"about_source": "file"})                       # no file chosen: the written text is used
    assert c["about_source"] == "text"
    c = app.sanitize({"about_source": "file", "about_file": "C:\\cv.docx"})
    assert c["about_source"] == "file" and c["about_file"] == "C:\\cv.docx"
    c = app.sanitize({"about_source": "file", "about_file": "\\\\server\\share\\cv.docx"})
    assert c["about_file"] == "" and c["about_source"] == "text"
    c = app.sanitize({"meeting_mode": "nope", "screen_delay": "9", "search_days": "5", "glossary": "x" * 9000,
                      "screen_provider": "../../evil"})
    assert c["meeting_mode"] == "general" and c["screen_delay"] == "3" and c["search_days"] == "90"
    assert len(c["glossary"]) == 3000 and c["screen_provider"] == ""
    assert "about_file" in app.INTERNAL_KEYS


# ---- About from a file --------------------------------------------------------------------------
def test_read_document_formats(app, tmp_path):
    d = str(tmp_path)
    p = os.path.join(d, "cv.txt")
    with open(p, "wb") as f:
        f.write("\ufeffJohn Doe\r\n\r\n\r\n\r\nNetwork engineer, 8 years. سلام".encode("utf-8"))
    t = app.read_document(p)
    assert t.startswith("John Doe") and "8 years" in t and "\r" not in t and "\n\n\n" not in t and "سلام" in t
    p = os.path.join(d, "cv.docx")
    with open(p, "wb") as f:
        f.write(app.docx_bytes([("title", "Jane Roe", None), ("p", "Skills: OSPF & <BGP>", None), ("p", "تجربه کاری", None)]))
    t = app.read_document(p)
    assert "Jane Roe" in t and "OSPF & <BGP>" in t and "تجربه کاری" in t
    try:
        import pypdf  # noqa: F401
    except ImportError:
        return
    p = os.path.join(d, "cv.pdf")
    _tiny_pdf(p, "Senior DBA at Example")
    assert "Senior DBA at Example" in app.read_document(p)


def test_read_document_errors(app, tmp_path):
    d = str(tmp_path)
    bad = app.DocError
    for path, why in (("", "No file"), (os.path.join(d, "missing.txt"), "not found"), ("\\\\srv\\x\\a.txt", "Network")):
        try:
            app.read_document(path)
            assert False, path
        except bad as e:
            assert why in str(e), (path, str(e))
    p = os.path.join(d, "a.exe")
    open(p, "wb").write(b"x")
    try:
        app.read_document(p)
        assert False
    except bad as e:
        assert ".docx" in str(e)
    p = os.path.join(d, "empty.txt")
    open(p, "w").write("  \n ")
    try:
        app.read_document(p)
        assert False
    except bad as e:
        assert "No text" in str(e)
    p = os.path.join(d, "broken.docx")
    open(p, "wb").write(b"not a zip")
    try:
        app.read_document(p)
        assert False
    except bad as e:
        assert "readable" in str(e)
    p = os.path.join(d, "broken.pdf")
    open(p, "wb").write(b"%PDF-1.4 nothing")
    try:
        app.read_document(p)
        assert False
    except bad:
        pass
    p = os.path.join(d, "long.txt")
    open(p, "w", encoding="utf-8").write("word " * 5000)
    assert len(app.read_document(p)) == app.ABOUT_MAX


def test_about_text_uses_file_and_falls_back(app, tmp_path):
    p = os.path.join(str(tmp_path), "me.txt")
    open(p, "w", encoding="utf-8").write("I am Siavash, 10 years in networking.")
    c = _cfg(app, about_me="written text", about_file=p, about_source="file")
    assert "Siavash" in app.about_text(c) and "written text" not in app.about_text(c)
    c["about_source"] = "text"
    assert app.about_text(c) == "written text"
    c["about_source"] = "file"
    os.remove(p)
    assert app.about_text(c) == "written text"                        # the file vanished: the written text is used
    info = app.about_info(c)
    assert info["file"] == "me.txt" and info["error"]


# ---- glossary and speakers -----------------------------------------------------------------------
def test_glossary_terms_and_spelling(app):
    c = _cfg(app, glossary="OSPF, RMAN\nkubectl; ospf ,  , k8s،Vlan 100")
    t = app.glossary_terms(c)
    assert t == ["OSPF", "RMAN", "kubectl", "k8s", "Vlan 100"]
    assert app.apply_glossary("we run ospf and rman, not Rmanager. kubectl get pods; K8S", t) == \
        "we run OSPF and RMAN, not Rmanager. kubectl get pods; k8s"
    assert app.glossary_terms(_cfg(app, glossary="")) == [] and app.glossary_note(_cfg(app)) == ""
    many = ",".join(f"t{i}" for i in range(200))
    assert len(app.glossary_terms(_cfg(app, glossary=many))) == app.GLOSSARY_MAX_TERMS


def test_glossary_reaches_the_prompts(app):
    class E(app.Engine):
        def __init__(self, cfg):
            self.cfg = cfg
    e = E(_cfg(app, glossary="OSPF, RMAN", context="Network job\nsecond line"))
    p = e._stt_prompt("them")
    assert p.startswith("Network job.") and "OSPF, RMAN" in p and len(p) <= 420
    assert E(_cfg(app, context="Network job\nx"))._stt_prompt("them") == "Network job"    # unchanged without a glossary


def test_speaker_label(app):
    assert app.speaker_label("Interviewer", None) == "Interviewer"
    assert app.speaker_label("Interviewer", 0) == "Interviewer"
    assert app.speaker_label("Interviewer", 1) == "Interviewer 2"
    assert app.speaker_label("Person 1", 2) == "Person 3"
    assert app.speaker_label("Person 1", True) == "Person 1"
    assert app.speaker_label("مصاحبه‌کننده", 1) == "مصاحبه‌کننده 2"


def test_deepgram_url_options(app):
    u = app.deepgram_url("nova-3", "en")
    assert "diarize" not in u and "keyterm" not in u
    u = app.deepgram_url("nova-3", "en", ["OSPF", "Vlan 100"], True)
    assert "diarize=true" in u and u.count("keyterm=") == 2 and "keyterm=Vlan+100" in u
    u = app.deepgram_url("nova-2", "en", ["OSPF"], False)
    assert "keywords=OSPF%3A2" in u and "keyterm" not in u


class _FakeWS:
    def __init__(self, msgs):
        self.msgs, self.closed = list(msgs), False

    def recv(self):
        if not self.msgs:
            self.closed = True
            return 0x8, b""
        return 0x1, json.dumps(self.msgs.pop(0)).encode()


def _dg_result(words, spk, final=True, speech_final=True):
    return {"type": "Results", "start": 1.0, "duration": 2.0, "is_final": final, "speech_final": speech_final,
            "channel": {"alternatives": [{"transcript": " ".join(w for w in words),
                                          "words": [{"word": w, "speaker": s} for w, s in zip(words, spk)]}]}}


def test_live_speaker_is_the_one_who_said_most_words(app):
    got = []
    lv = app.DeepgramLive("them", "k", "nova-3", "en", "", lambda *a: None, lambda *a: got.append(a), lambda e: None,
                          diarize=True)
    lv.stop_event.set()
    lv.ws = _FakeWS([_dg_result(["hello", "there", "all"], [1, 1, 0])])
    lv._receive()
    assert len(got) == 1 and got[0][3] == "hello there all" and got[0][4] == 1
    got.clear()
    lv2 = app.DeepgramLive("them", "k", "nova-3", "en", "", lambda *a: None, lambda *a: got.append(a), lambda e: None)
    lv2.stop_event.set()
    lv2.ws = _FakeWS([_dg_result(["hello", "there"], [1, 1])])
    lv2._receive()
    assert len(got) == 1 and len(got[0]) == 4                          # no speaker argument when it is switched off


def test_session_keeps_speaker_feedback_and_screens(app):
    c = _cfg(app, them_label="Interviewer")
    s = app.Session(c)
    a = s.add("them", 100.0, 102.0, "First question?", "en", [], 0)
    b = s.add("them", 103.0, 105.0, "Second person here", "en", [], 1)
    assert a["speaker"] is None and b["speaker"] == 1
    assert s.label(s.get(b["id"])) == "Interviewer 2" and s.label(s.get(a["id"])) == "Interviewer"
    s.set_feedback("## Overall\nGood.")
    s.add_screen("The screen shows a form.")
    s.save_if_dirty()
    md = open(s.path, encoding="utf-8").read()
    assert "Interviewer 2:** Second person here" in md and "## Feedback" in md and "## Screen" in md
    r = app.Session.load_last(c)
    assert r is not None and r.feedback == "## Overall\nGood." and r.screens[0]["text"] == "The screen shows a form."
    assert sorted(x["speaker"] or 0 for x in r.entries.values()) == [0, 1]


# ---- statistics ------------------------------------------------------------------------------------
def test_talk_stats(app):
    rows = [{"source": "them", "t0": 0, "t_end": 5, "text": "Tell me about OSPF please"},
            {"source": "me", "t0": 13, "t_end": 43, "text": "Um well " + "word " * 58 + "you know"},
            {"source": "them", "t0": 44, "t_end": 46, "text": "Thanks"},
            {"source": "me", "t0": 47, "t_end": 49, "text": "Sure"}]
    st = app.talk_stats(rows)
    assert st["me_words"] == 63 and st["them_words"] == 6 and st["me_share"] == 91
    assert st["wpm"] == 118 and st["fillers"] == 2 and st["long_pauses"] == 1 and st["answers"] == 2
    assert st["avg_delay"] == 4.5
    assert app.talk_stats([])["me_share"] == 0 and app.talk_stats([])["wpm"] is None


# ---- screen picture ----------------------------------------------------------------------------------
def _read_png(data):
    assert data[:8] == b"\x89PNG\r\n\x1a\n"
    pos, idat, ihdr = 8, b"", None
    while pos < len(data):
        n, kind = struct.unpack(">I4s", data[pos:pos + 8])
        body = data[pos + 8:pos + 8 + n]
        assert struct.unpack(">I", data[pos + 8 + n:pos + 12 + n])[0] == zlib.crc32(kind + body) & 0xFFFFFFFF
        if kind == b"IHDR":
            ihdr = struct.unpack(">IIBBBBB", body)
        elif kind == b"IDAT":
            idat += body
        pos += 12 + n
    w, h = ihdr[0], ihdr[1]
    raw = zlib.decompress(idat)
    assert len(raw) == h * (w * 3 + 1)
    return w, h, raw


def test_png_encoder_and_downscale(app):
    import numpy as np
    w, h = 300, 100
    px = np.zeros((h, w, 4), dtype=np.uint8)
    px[:, :, 0], px[:, :, 1], px[:, :, 2] = 10, 20, 200            # BGRA: blue 10, green 20, red 200
    ww, hh, raw = _read_png(app.encode_png(w, h, px.tobytes()))
    assert (ww, hh) == (300, 100) and list(raw[1:4]) == [200, 20, 10]
    big = np.zeros((1000, 4000, 4), dtype=np.uint8)
    big[:, :, 2] = 255
    ww, hh, raw = _read_png(app.encode_png(4000, 1000, big.tobytes(), 1920))
    assert (ww, hh) == (1920, 480) and list(raw[1:4]) == [255, 0, 0]


def test_vision_model_guess(app):
    assert app.guess_vision_model(["whisper-1", "text-embedding-3-small", "gpt-4o-mini", "gpt-4o"]) == "gpt-4o-mini"
    assert app.guess_vision_model(["llama-3.3-70b", "meta-llama/llama-4-scout-17b-16e-instruct"]) == \
        "meta-llama/llama-4-scout-17b-16e-instruct"
    assert app.guess_vision_model(["llama-3.3-70b", "whisper-large-v3"]) is None
    assert app.guess_vision_model([]) is None


def test_picture_messages_are_not_counted_as_text(app):
    m = {"role": "user", "content": [{"type": "text", "text": "abcd"},
                                     {"type": "image_url", "image_url": {"url": "data:image/png;base64," + "A" * 900000}}]}
    assert app._msg_chars(m) < 10000 and app._msg_chars({"content": "hello"}) == 5


# ---- old meetings ----------------------------------------------------------------------------------------
def _meeting(app, when, lines):
    s = app.Session(_cfg(app))
    s.started = when
    s.path = os.path.join(app.MEETINGS_DIR, when.strftime("meeting_%Y-%m-%d_%H-%M-%S.md"))
    t = when.timestamp()
    for i, (src, text, tr) in enumerate(lines):
        e = s.add(src, t + i * 10, t + i * 10 + 5, text, "en", [])
        s.update(e["id"], translation=tr)
    s.save_if_dirty()
    return s


def test_search_old_meetings_with_a_range(app):
    now = datetime.datetime(2026, 9, 29, 12, 0)
    _meeting(app, now - datetime.timedelta(days=3), [("them", "Tell me about OSPF areas", "درباره OSPF بگو"),
                                                     ("me", "Area zero is the backbone", "")])
    _meeting(app, now - datetime.timedelta(days=60), [("them", "What is BGP?", ""), ("me", "OSPF is not BGP", "")])
    _meeting(app, now - datetime.timedelta(days=400), [("them", "Old talk about OSPF", "")])
    hits, n = app.search_meetings(app.MEETINGS_DIR, "ospf", 7, now=now)
    assert n == 1 and len(hits) == 1 and hits[0]["date"] == "2026-09-26 12:00" and "درباره" in hits[0]["text"]
    hits, n = app.search_meetings(app.MEETINGS_DIR, "OSPF", 90, now=now)
    assert n == 2 and len(hits) == 2
    hits, n = app.search_meetings(app.MEETINGS_DIR, "OSPF", 0, now=now)
    assert n == 3 and len(hits) == 3
    hits, _ = app.search_meetings(app.MEETINGS_DIR, "backbone area", 0, now=now)     # all words must be in one line
    assert len(hits) == 1 and hits[0]["who"]
    hits, _ = app.search_meetings(app.MEETINGS_DIR, "backbone bananas", 0, now=now)   # half of the words is enough
    assert len(hits) == 1
    assert app.search_meetings(app.MEETINGS_DIR, "?", 0, now=now) == ([], 0)
    assert app.fold("كيك") == app.fold("کیک") and app.fold("Café") == "cafe"


def test_search_survives_odd_files(app):
    os.makedirs(app.MEETINGS_DIR, exist_ok=True)
    open(os.path.join(app.MEETINGS_DIR, "meeting_2026-09-01_10-00-00.md"), "wb").write(b"\xff\xfe**[10:00:00] Me:** bad bytes ospf\n")
    open(os.path.join(app.MEETINGS_DIR, "notes.md"), "w").write("ospf")
    open(os.path.join(app.MEETINGS_DIR, "meeting_broken.md"), "w").write("ospf")
    hits, n = app.search_meetings(app.MEETINGS_DIR, "ospf", 0, now=datetime.datetime(2026, 9, 29))
    assert n == 1


# ---- export ---------------------------------------------------------------------------------------------
def _sample_session(app):
    c = _cfg(app, context="Interview", them_label="Interviewer")
    s = app.Session(c)
    t = s.started.timestamp()
    e = s.add("them", t + 5, t + 8, "What is <OSPF> & why?", "en", [])
    s.update(e["id"], translation="OSPF چیست؟", answer="It is a link-state protocol.", answer_fa="یک پروتکل است.")
    s.add("me", t + 12, t + 14, "Let me explain.", "en", [])
    s.set_summary("## Summary\n- one\n- two")
    return s


def test_export_docx_is_a_valid_word_file(app):
    s = _sample_session(app)
    data = app.export_bytes(s, "docx")
    z = zipfile.ZipFile(io.BytesIO(data))
    assert z.testzip() is None and "word/document.xml" in z.namelist()
    import xml.dom.minidom
    xml.dom.minidom.parseString(z.read("word/document.xml"))
    try:
        import docx
    except ImportError:
        return
    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "What is <OSPF> & why?" in text and "OSPF چیست؟" in text and "• one" in text
    assert "Interviewer:" in text
    rtl = [p for p in d.paragraphs if "چیست" in p.text][0]
    assert rtl._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is not None


def test_export_srt_txt_json(app):
    s = _sample_session(app)
    srt = app.export_bytes(s, "srt").decode("utf-8")
    assert srt.startswith("1\n00:00:05,000 --> 00:00:08,000\nInterviewer: What is <OSPF> & why?\nOSPF چیست؟\n")
    assert "2\n00:00:12,000 --> 00:00:14,000\n" in srt
    txt = app.export_bytes(s, "txt").decode("utf-8")
    assert "Interviewer: What is" in txt and "== Summary ==" in txt
    assert json.loads(app.export_bytes(s, "json"))["entries"][0]["text"].startswith("What is")
    assert app.srt_time(3661.5) == "01:01:01,500" and app.srt_time(-4) == "00:00:00,000"


# ---- hidden window -----------------------------------------------------------------------------------------
class _Stop:
    """A stop flag whose wait() also advances a fake clock, so the keeper loop can be run without sleeping."""
    def __init__(self, ticks):
        self.ticks = ticks
        self.flag = False

    def is_set(self):
        return self.flag

    def wait(self, t=None):
        self.ticks -= 1
        if self.ticks <= 0:
            self.flag = True
        return self.flag


def test_hiding_is_set_and_reported(app):
    log = []
    state = {"aff": 0}
    set_aff = lambda h, f: state.update(aff=f) or True
    ok = app.keep_hidden(lambda: 77, set_aff, lambda h: state["aff"], lambda *a: log.append(a), _Stop(3), interval=0)
    assert ok and log[0] == (True, "exclude", "") and state["aff"] == 0x11


def test_hiding_falls_back_to_black_and_notices_a_loss(app):
    log = []
    state = {"aff": 0, "refuse_exclude": True, "sets": 0}

    def set_aff(h, f):
        state["sets"] += 1
        if f == 0x11 and state["refuse_exclude"]:
            return False
        state["aff"] = f
        return True

    class Stop(_Stop):
        def wait(self, t=None):
            if self.ticks == 3:
                state["aff"] = 0                        # something removed the flag
            return super().wait(t)
    ok = app.keep_hidden(lambda: 5, set_aff, lambda h: state["aff"], lambda *a: log.append(a), Stop(4), interval=0)
    assert ok and log[0] == (True, "black", "")
    assert state["aff"] == 1                             # it was put back
    assert log[-1][0] is True


def test_hiding_failure_is_reported_not_hidden(app):
    log = []
    assert not app.keep_hidden(lambda: 5, lambda h, f: False, lambda h: 0, lambda *a: log.append(a), _Stop(3), interval=0)
    assert log[-1][0] is False and "refused" in log[-1][2]
    log.clear()
    assert not app.keep_hidden(lambda: None, lambda h, f: True, lambda h: 0, lambda *a: log.append(a), _Stop(1), first_wait=0)
    assert log[-1][0] is False and "not be found" in log[-1][2]


def test_hide_status_file(app, tmp_path):
    p = os.path.join(str(tmp_path), "hide.json")
    assert app.read_hide_status(p) is None
    app.write_hide_status(p, True, "exclude", "")
    assert app.read_hide_status(p) == {"active": True, "mode": "exclude", "error": "", "pid": os.getpid()}
    open(p, "w").write("{broken")
    assert app.read_hide_status(p) is None
    open(p, "w").write("[1]")
    assert app.read_hide_status(p) is None


def test_hide_setting_is_read_from_the_settings_file(app):
    assert app.hide_wanted() is False
    with open(app.CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump({"hide_from_share": True}, f)
    assert app.hide_wanted() is True
    with open(app.CONFIG_PATH, "w", encoding="utf-8") as f:
        f.write("not json")
    assert app.hide_wanted() is False
    if sys.platform != "win32":
        assert app.hidden_window_possible() is False
    else:
        assert isinstance(app.hidden_window_possible(), bool)   # true when pywebview is installed


def test_glossary_short_mixed_case_left_alone(app):
    assert app.apply_glossary("let's go now", ["Go"]) == "let's go now"
    assert app.apply_glossary("use ospf", ["OSPF"]) == "use OSPF"


def test_search_question_style_and_order(app, tmp_path):
    import pathlib; tmp_path = pathlib.Path(tmp_path)
    import datetime
    (tmp_path / "meeting_2026-01-05_10-00-00.md").write_text("**[00:00:01] Interviewer:** we discussed kubernetes rollout\n", encoding="utf-8")
    (tmp_path / "recording_x_2025-12-01_10-00-00.md").write_text("**[00:00:01] Interviewer:** unrelated\n", encoding="utf-8")
    hits, n = app.search_meetings(str(tmp_path), "what did we say about kubernetes?", days=0)
    assert n == 2 and len(hits) == 1 and hits[0]["file"].startswith("meeting_")


def test_srt_base_for_recording(app):
    import datetime
    d = datetime.datetime(2026, 1, 5, 15, 30)
    base = datetime.datetime.combine(d.date(), datetime.time()).timestamp()
    rows = [{"text": "hi", "t0": base + 5, "t_end": base + 7}]
    out = app.srt_text(rows, base, lambda r: "X")
    assert "00:00:05,000 --> 00:00:07,000" in out


def test_decode_text_encodings(app):
    assert app._decode_text("café".encode("cp1252")) == "café"
    fa = "سلام دنيا اين يك متن است"
    assert app._decode_text(fa.encode("cp1256")) == fa
    assert app._decode_text(fa.encode("utf-16")) == fa
    assert app._decode_text(fa.encode("utf-8")) == fa


def test_glossary_two_letter_terms(app):
    assert app.apply_glossary("I think it is fine", ["IT", "AI"]) == "I think it is fine"


def test_docx_text_no_duplicates(app, tmp_path):
    import zipfile, os
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    x = (f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>Hello</w:t></w:r></w:p>'
         f'<w:p><w:r><w:txbxContent><w:p><w:r><w:t>Box</w:t></w:r></w:p></w:txbxContent></w:r></w:p></w:body></w:document>')
    p = os.path.join(tmp_path, "a.docx")
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("word/document.xml", x)
    assert app._docx_text(p).split("\n") == ["Hello", "Box"]


def test_model_info_and_picks(app):
    raw = [{"id": "whisper-large-v3", "active": True}, {"id": "old-model", "active": False},
           {"id": "qwen/qwen3.8-27b", "active": True, "input_modalities": ["text", "image"]},
           {"id": "openai/gpt-oss-120b", "active": True}, {"id": "llama-guard", "active": True},
           {"id": "text-embedding-3", "active": True}]
    infos = app.model_infos(raw)
    by = {i["id"]: i for i in infos}
    assert by["whisper-large-v3"]["stt"] and not by["whisper-large-v3"]["chat"]
    assert not by["old-model"]["alive"]
    assert by["qwen/qwen3.8-27b"]["vision"] and by["qwen/qwen3.8-27b"]["sure"]
    assert not by["llama-guard"]["chat"] and not by["text-embedding-3"]["chat"]
    assert app.pick_vision(infos) == "qwen/qwen3.8-27b"
    assert app.pick_vision([{"id": "x", "vision": True, "alive": False, "sure": True, "chat": True, "stt": False}]) is None
    assert app.pick_chat(infos)[0] == "openai/gpt-oss-120b"
    # a model the service says cannot read pictures is not chosen even if its name looks like one
    assert app.model_info({"id": "gpt-4o-audio", "input_modalities": ["text"]})["vision"] is False
    # the id is kept exactly as the service writes it (it is what requests use)
    assert app.model_info({"id": "models/gemini-2.5-flash"})["id"] == "models/gemini-2.5-flash"
    assert not app.model_info({"id": "flux", "type": "image"})["chat"]
    assert not app.model_info({"id": "x-old", "deprecation": "2026-01-01"})["alive"]
    assert app.pick_chat(app.model_infos([{"id": "openai/gpt-oss-120b"}, {"id": "acme-128b-x"}, {"id": "groq/compound"}]))[0] == "openai/gpt-oss-120b"


def test_custom_item_and_needs(app):
    it = app.custom_item("llm", "unsloth/gemma-3-4b-it-GGUF", "gemma-3-4b-it-Q4_K_M.gguf", 2490)
    assert it["file"].endswith(".gguf") and "/" not in it["id"] and it["id"].startswith("hf-")
    it2 = app.custom_item("stt", "Systran/faster-whisper-small", "", 484)
    assert "file" not in it2
    for bad in (("llm", "../x/y", "a.gguf", 5), ("llm", "a/b", "../a.gguf", 5), ("llm", "a/b", "a.bin", 5), ("stt", "a/b", "", 0), ("stt", "bad", "", 5)):
        try:
            app.custom_item(*bad); assert False, bad
        except ValueError:
            pass
    n = app.model_needs("llm", 2490)
    assert n["ram_mb"] > 2490 and "cores" in n["cpu"]
    assert app.model_fit(2000, 16000) == "ok" and app.model_fit(10000, 16000) == "tight" and app.model_fit(15000, 16000) == "no"
    pick = app._hf_pick_file([{"rfilename": "m-Q8_0.gguf", "size": 8_000_000_000}, {"rfilename": "m-Q4_K_M.gguf", "size": 4_000_000_000},
                              {"rfilename": "m-00001-of-00002.gguf", "size": 1}, {"rfilename": "mmproj-x.gguf", "size": 5}])
    assert pick == ("m-Q4_K_M.gguf", 4000)


def test_decode_partial_utf8_and_srt_blank_lines(app):
    fa = "سلام دنیا " * 50
    raw = (fa * 3000).encode("utf-8")[:1_500_001]        # cut in the middle of a character
    assert app._decode_text(raw).startswith("سلام دنیا")
    rows = [{"text": "a\n\nb", "translation": "x\n\n\ny", "t0": 1, "t_end": 3}, {"text": "c", "t0": 4, "t_end": 5}]
    blocks = app.srt_text(rows, 0, lambda r: "L").strip().split("\n\n")
    assert len(blocks) == 2 and blocks[0].split("\n")[2:] == ["a b", "x y"]   # no blank line inside a cue



def test_sanitize_float_choice_and_hf_search(app):
    import httpx, json
    cfg = app.sanitize({"search_days": 30.0, "screen_delay": 2.0})
    assert cfg["search_days"] == "30" and cfg["screen_delay"] == "2"
    a = app.custom_item("llm", "Org/Model", "m-Q4_K_M.gguf", 100)
    b = app.custom_item("llm", "org/model", "m-q4_k_m.gguf", 100)
    c = app.custom_item("llm", "Org/Model", "m-Q8_0.gguf", 100)
    assert a["id"] == b["id"] and a["id"] != c["id"] and len(a["id"]) < 60      # same folder only for the same file

    real = httpx.Client

    def make(handler):
        app.httpx.Client = lambda **kw: real(transport=httpx.MockTransport(handler), **{k: v for k, v in kw.items() if k != "proxy"})
    make(lambda req: httpx.Response(200, json={"error": "nope"}))
    try:
        app.hf_search("llm", "popular", ""); assert False
    except app.APIError:
        pass

    def h(req):
        if req.url.path == "/api/models":
            return httpx.Response(200, json=[{"id": "a/b-Instruct-GGUF", "downloads": "x"}, "junk",
                                             {"modelId": "a/c-Instruct-GGUF", "downloads": 5000}])
        return httpx.Response(200, json={"siblings": ["bad", {"rfilename": "c-Q4_K_M.gguf", "size": "not-a-number"},
                                                      {"rfilename": "c-Q5_K_M.gguf", "size": 5_000_000_000}]})
    make(h)
    res = app.hf_search("llm", "popular", "")
    assert [i["file"] for i in res["items"]] == ["c-Q5_K_M.gguf"] and res["items"][0]["mb"] == 5000
    try:
        make(lambda req: httpx.Response(200, json=[{"id": "a/c-Instruct-GGUF", "downloads": 5000}]) if req.url.path == "/api/models" else httpx.Response(429))
        try:
            app.hf_search("llm", "popular", ""); assert False
        except app.APIError as e:
            assert "429" in str(e)
    finally:
        app.httpx.Client = real


def test_hide_status_stale_and_ping_empty(app, tmp_path):
    import json, os, time
    p = os.path.join(tmp_path, "h.json")
    json.dump({"active": True, "mode": "exclude", "error": "", "t": time.time() - 500}, open(p, "w"))
    st = app.read_hide_status(p)
    assert st["active"] is False and "stopped" in st["error"]
    json.dump({"active": True, "mode": "exclude", "error": "", "t": time.time()}, open(p, "w"))
    assert app.read_hide_status(p)["active"] is True
    # an empty or odd model list is not an error (6.1 behaviour)
    import httpx
    api = app.GroqAPI.__new__(app.GroqAPI)
    for body in ({"data": []}, [], ["gpt-a", "gpt-b"], {"models": [{"id": "m1"}]}):
        api.request = lambda *a, _b=body, **k: httpx.Response(200, json=_b, request=httpx.Request("GET", "http://x/models"))
        ms, ids = api.ping()
        assert isinstance(ids, list)
    assert ids == ["m1"]


def _speech_events(app, sens, level, seconds=1.0, gain=1.0):
    """Feeds a steady tone of the given loudness to a Segmenter and returns what it reported."""
    import numpy as np
    events = []
    seg = app.Segmenter("me", 16000, sens, lambda kind, *a, **k: events.append(kind))
    n = int(16000 * app.FRAME_SEC)
    t = np.arange(n) / 16000.0
    frame = (np.sin(2 * np.pi * 220 * t) * level * 1.414 * gain).astype("float32")
    silent = np.zeros(n, dtype="float32")
    room = (np.random.RandomState(2).randn(n) * 0.0004).astype("float32")
    for _ in range(int(3.0 / app.FRAME_SEC)):                 # a quiet room first (sets the noise floor)
        seg.feed(room)
    for _ in range(int(seconds / app.FRAME_SEC)):
        seg.feed(frame)
    for _ in range(int(2.0 / app.FRAME_SEC)):
        seg.feed(silent)
    return events


def test_low_sensitivity_ignores_small_sounds(app):
    assert "segment" in _speech_events(app, 6, 0.02)          # normal speech level: heard at the default
    assert "segment" not in _speech_events(app, 1, 0.02)      # the same sound is ignored at the lowest setting
    assert "segment" in _speech_events(app, 1, 0.12)          # loud speech is still heard
    assert "segment" not in _speech_events(app, 2, 0.05, seconds=0.2)   # a short click is ignored


def test_sensitivity_scale_is_ordered(app):
    lv = []
    for s in range(1, 11):
        seg = app.Segmenter("me", 16000, s, lambda *a, **k: None)
        lv.append(seg.min_level)
    assert lv == sorted(lv, reverse=True)
    assert abs(lv[5] - 0.0056) < 0.0005                        # the default is unchanged


def test_mic_gain_is_limited_and_only_for_me(app):
    me = app.AudioCapture(None, {}, "me", 6, lambda *a, **k: None)
    them = app.AudioCapture(None, {}, "them", 6, lambda *a, **k: None)
    me.set_gain(1000)
    them.set_gain(300)
    assert me.gain == 3.0 and them.gain == 1.0
    me.set_gain(1)
    assert me.gain == 0.25
    assert app.sanitize({"mic_gain": 5000})["mic_gain"] == 300


def test_live_service_hears_the_mic_only_during_speech(app):
    import numpy as np
    got = []
    cap = app.AudioCapture(None, {}, "me", 1, lambda *a, **k: None)
    cap.seg = app.Segmenter("me", 16000, 1, lambda *a, **k: None)
    cap.tap = lambda x, rate: got.append(float(np.abs(x).max()) if x is not None else 0.0)
    quiet = (np.random.RandomState(1).randn(480) * 0.004).astype("float32")
    for _ in range(20):
        cap._got(quiet, 16000)
    assert got and max(got) == 0.0                              # a quiet room is sent as silence


def test_looks_askable(app):
    ask = ["Tell me about yourself.", "Walk me through your last project", "How does BGP choose a path",
           "and why?", "Can you give an example", "Describe how you would design a backup plan.",
           "Erzählen Sie mir von Ihrer Erfahrung", "Was ist der Unterschied zwischen TCP und UDP"]
    quiet = ["We are a company founded in 2010 in Berlin.", "Great, thank you.", "Our team has twelve people and we ship monthly."]
    for s in ask:
        assert app.looks_askable(s, "en"), s
    for s in quiet:
        assert not app.looks_askable(s, "en"), s
    assert app.looks_askable("Please introduce yourself", "en") and app.looks_askable("Walk us through your CV", "en")
    assert app.looks_askable("سلام معرفی کنید", "")                # unknown language, not plain English: never dropped
    assert app.looks_askable("سلام خوش آمدید", "fa")           # other languages: always yes


class _Hub:
    def publish(self, *a, **k): pass
    def toast(self, *a, **k): pass


class _FakeApp:
    def __init__(self, app):
        self.cfg = app.sanitize({})
        self.hub = _Hub()
        self.stats = app.Stats()
        self.captures = []
        self.running = False
        self.engine = None
    def provider(self, pid): return None


def _engine(app):
    fa = _FakeApp(app)
    sess = app.Session(fa.cfg)
    eng = app.Engine(fa, sess, workers=False)
    calls = []
    eng._answer = lambda eid, force=False, question=None, _retry=0: calls.append((eid, force, question))
    eng._submit_tr = lambda eid: True
    return eng, sess, calls


def test_no_live_translation_when_the_meeting_is_in_my_language(app):
    eng, sess, calls = _engine(app)
    sent = []
    eng.pv_pool.submit = lambda fn, *a: sent.append(a)
    try:
        eng.cfg["tr_provider"] = "groq"
        eng.cfg["my_language"] = "fa"
        eng.cfg["languages"] = ["fa"]                           # a Persian-only meeting, my language Persian
        eng._preview_translate("k1", "them", "سلام این یک جمله‌ی نیمه‌تمام است", min_new=1)
        assert sent == []                                       # nothing is sent to be translated
        eng.cfg["languages"] = ["en"]                           # an English meeting: the live translation still runs
        eng._preview_translate("k2", "them", "this is an unfinished sentence", min_new=1)
        assert len(sent) == 1
        eng.cfg["languages"] = ["en", "fa"]                     # two languages: the language is not known yet, so it runs
        eng._preview_translate("k3", "them", "this is another unfinished sentence", min_new=1)
        assert len(sent) == 2
    finally:
        eng.stop_event.set()


def test_gguf_arch_tells_speech_from_chat_and_routes_each(app, tmp_path):
    import struct
    def gguf(name, arch):
        path = str(tmp_path / name); key = b"general.architecture"; val = arch.encode()
        with open(path, "wb") as f:
            f.write(b"GGUF" + struct.pack("<IQQ", 3, 0, 1)
                    + struct.pack("<Q", len(key)) + key + struct.pack("<I", 8)
                    + struct.pack("<Q", len(val)) + val)
        return path
    wh = gguf("whisper-medium-Q8_0.gguf", "whisper")
    gem = gguf("gemma3-4b.gguf", "gemma3")
    assert app.gguf_arch(wh) == "whisper"
    assert app.gguf_arch(gem) == "gemma3"
    assert app.gguf_arch(str(tmp_path)) == ""                      # a folder, not a .gguf
    # a Whisper model does not belong in the Local AI model slot
    assert "speech model" in app.wrong_slot_message(wh, "llm").lower()
    assert app.wrong_slot_message(gem, "llm") == ""
    # a chat model does not belong in the speech-to-text slot
    assert "chat model" in app.wrong_slot_message(gem, "stt").lower()
    assert app.wrong_slot_message(wh, "stt") == ""
    # the speech picker now points a chat model to the right place instead of "not a speech model"
    _, err = app.model_folder(gem)
    assert "Local AI model" in err


def _say(app, eng, text, when=None):
    now = when or __import__("time").time()
    eng.accept_text("them", text, "en", now - 2, now - 0.1, [next(app._seg_counter)],
                    {"pause": 0.0, "queue": 0.0, "stt": 0.0, "service": "x", "model": "x"})


def test_question_in_pieces_gets_one_answer(app):
    import time
    eng, sess, calls = _engine(app)
    try:
        _say(app, eng, "Tell me about a time")
        time.sleep(0.4)
        _say(app, eng, "when you led a team through an outage")
        time.sleep(1.8)
        assert len(calls) == 1, calls
        assert "Tell me about a time" in calls[0][2] and "outage" in calls[0][2]
        first = sess.ordered()[0]
        assert first["ans_state"] == "none"                     # the first piece was merged into the second
    finally:
        eng.stop_event.set()


def test_statements_and_old_lines_are_not_answered(app):
    import time
    eng, sess, calls = _engine(app)
    try:
        _say(app, eng, "We are a company founded in 2010 in Berlin.")
        _say(app, eng, "Great, thank you.")
        eng.accept_text("them", "How would you monitor a database?", "en", time.time() - 200, time.time() - 199, [next(app._seg_counter)],
                        {"pause": 0.0, "queue": 0.0, "stt": 0.0, "service": "x", "model": "x"})
        time.sleep(1.6)
        assert calls == []
    finally:
        eng.stop_event.set()


def test_question_mark_is_answered_quickly(app):
    import time
    eng, sess, calls = _engine(app)
    try:
        _say(app, eng, "How would you monitor a database?")
        time.sleep(0.7)
        assert len(calls) == 1 and calls[0][2] is None
    finally:
        eng.stop_event.set()


def test_answer_button_cancels_a_waiting_automatic_answer(app):
    import time
    eng, sess, calls = _engine(app)
    try:
        _say(app, eng, "Explain how OSPF elects a DR")
        eng.answer_now()
        time.sleep(1.6)
        assert [c[1] for c in calls] == [True]                  # only the forced one
    finally:
        eng.stop_event.set()


def test_stopped_meeting_leaves_no_answer_writing(app):
    import time
    eng, sess, calls = _engine(app)
    _say(app, eng, "Explain how OSPF elects a DR")
    eng.stop_event.set()                                        # the meeting ends during the short wait
    time.sleep(1.6)
    assert calls == []
    assert sess.ordered()[0]["ans_state"] == "none"


def test_retry_does_not_override_a_forced_answer(app):
    eng, sess, calls = _engine(app)
    try:
        _say(app, eng, "How would you monitor a database?")
        eid = sess.ordered()[0]["id"]
        eng.ans_gen[eid] += 1                                   # the user pressed F2 meanwhile (a newer request)
        eng._retry_answer(eid, None, 1, 0)                      # the old retry wakes up
        assert calls == []
    finally:
        eng.stop_event.set()


def test_repeated_question_is_found(app):
    import time
    eng, sess, calls = _engine(app)
    now = time.time()
    a = sess.add("them", now - 200, now - 195, "How would you design a backup strategy for a large Oracle database?", "en", [], None)
    sess.add("me", now - 190, now - 150, "I would use RMAN and take backups every night, that is it.", "en", [], None)
    b = sess.add("them", now - 20, now - 15, "So how would you design a backup strategy for a big Oracle database?", "en", [], None)
    n, prev, said = eng.find_repeats(b, None)
    assert n == 1 and prev["id"] == a["id"] and "RMAN" in said
    c = sess.add("them", now - 5, now - 2, "Tell me about your experience with OSPF and BGP routing", "en", [], None)
    assert eng.find_repeats(c, None)[0] == 0
    # the same question in two pieces (no answer between) is not "asked again"
    d1 = sess.add("them", now + 1, now + 3, "How would you monitor replication lag on a standby", "en", [], None)
    d2 = sess.add("them", now + 4, now + 6, "How would you monitor replication lag on a standby database", "en", [], None)
    assert eng.find_repeats(d2, None)[0] == 0


def test_repeat_is_never_skipped_and_reaches_the_prompt(app):
    import time
    eng, sess, calls = _engine(app)
    eng.__dict__.pop("_answer", None)                            # use the real _answer
    seen = {}

    def fake_chat(chain, messages, max_tokens, temperature, on_text):
        seen["user"] = messages[-1]["content"]
        seen["system"] = messages[0]["content"]
        return "NO_REPLY"
    eng.run_chat = fake_chat
    eng.chat_targets = lambda task: [("groq", "m", None)]
    now = time.time()
    sess.add("them", now - 200, now - 195, "How would you design a backup strategy for a large Oracle database?", "en", [], None)
    sess.add("me", now - 190, now - 150, "I would use RMAN nightly.", "en", [], None)
    b = sess.add("them", now - 20, now - 15, "How would you design a backup strategy for a large Oracle database?", "en", [], None)
    eng._answer(b["id"])
    assert "REPEATED QUESTION" in seen["user"] and "RMAN nightly" in seen["user"]
    assert sess.get(b["id"])["repeat"] == 1
    assert sess.get(b["id"])["ans_state"] == "done"              # NO_REPLY is not accepted for an asked-again question


def test_situation_card_reads_the_model_answer(app):
    eng, sess, calls = _engine(app)
    import time
    now = time.time()
    sess.add("them", now - 50, now - 45, "Tell me about yourself", "en", [], None)
    sess.add("me", now - 40, now - 20, "I am a network engineer.", "en", [], None)
    sess.add("them", now - 10, now - 5, "How does OSPF elect a designated router", "en", [], None)
    published = []
    eng.hub = type("H", (), {"publish": lambda self, *a, **k: published.append((a, k)), "toast": lambda *a, **k: None})()
    eng.chat_targets = lambda task: [("groq", "m", None)]
    eng.run_chat = lambda *a, **k: 'Sure: {"phase": "technical", "topic": "OSPF DR election", "difficulty": 7, "trend": "up", "signal": "meh", "tip": "Answer directly."}'
    eng._coach_run()
    st = eng.coach_state
    assert st["phase"] == "technical" and st["difficulty"] == 5 and st["trend"] == "up" and st["signal"] == "neutral"
    assert published and published[0][0][0] == "coach"
    eng.run_chat = lambda *a, **k: "not json at all"
    eng._coach_run()                                             # a bad reply changes nothing and raises nothing
    assert eng.coach_state is st


def test_similar_frames_are_not_repeats(app):
    sim = app.Engine._sim
    assert sim("How would you back up the production database?", "How would you restore the production database?") < app.REPEAT_SIM
    assert sim("Tell me about your experience with Oracle", "Tell me about your experience with Kubernetes") < app.REPEAT_SIM
    assert sim("Can you walk me through how you would design the backup?", "Can you walk me through how you would test the backup?") < app.REPEAT_SIM
    assert sim("How would you design a backup strategy for a large Oracle database?",
               "So how would you design a backup strategy for a big Oracle database?") >= app.REPEAT_SIM
    assert sim("Why did you leave your last job", "Why did you leave your last job") >= app.REPEAT_SIM


def test_coach_reply_parsing_is_forgiving(app):
    import time
    eng, sess, calls = _engine(app)
    now = time.time()
    sess.add("them", now - 50, now - 45, "Tell me about yourself", "en", [], None)
    sess.add("them", now - 10, now - 5, "How does OSPF elect a designated router", "en", [], None)
    eng.chat_targets = lambda task: [("groq", "a", None), ("groq", "b", None)]
    used = []
    eng.hub = type("H", (), {"publish": lambda *a, **k: None, "toast": lambda *a, **k: None})()

    def chat(chain, *a, **k):
        used.append(chain)
        return '```json\n{"phase": "intro", "topic": "x", "difficulty": 2} trailing } text\n```'
    eng.run_chat = chat
    eng._coach_run()
    assert eng.coach_state["phase"] == "intro"
    assert used[0] == [("groq", "b", None)]                     # its own model, not the first one the answers use
    eng.run_chat = lambda *a, **k: "[1, 2]"
    eng._coach_run()                                             # a list is not a situation
    assert eng.coach_state["phase"] == "intro"


def test_question_types_and_targets(app):
    cases = {"Tell me about yourself": "intro", "Tell me about a time you disagreed with your manager": "behavioural",
             "How would you design a backup strategy for a large database?": "design",
             "What would you do if the database suddenly becomes slow?": "troubleshoot",
             "What is the difference between TCP and UDP?": "concept",
             "Do you have any questions for us?": "candidate_q", "What are your salary expectations?": "salary",
             "What is your notice period?": "availability", "What is your greatest weakness?": "weakness",
             "Have you worked with Oracle Data Guard?": "experience", "Can you write a function that reverses a list?": "coding",
             "Why do you want to work here?": "motivation", "Do you know Linux?": "yesno"}
    for text, want in cases.items():
        got = app.classify_question(text)[0]
        assert got == want, (text, got, want)
    assert app.classify_question("We are a company founded in 2010")[0] == ""
    name, lo, hi = app.classify_question("Tell me about yourself")
    assert (lo, hi) == (60, 90) and app.qtype_hint("intro", "fa") and app.qtype_hint("intro", "en") != app.qtype_hint("intro", "fa")


def _rows(app, spec):
    """spec: [(source, t0, t_end, text, extra)] -> entries with ids"""
    out = []
    for i, (src, a, b, text, extra) in enumerate(spec, 1):
        out.append({"id": i, "source": src, "t0": a, "t_end": b, "text": text, **extra})
    return out


def test_reply_check_flags(app):
    q = {"id": 1, "source": "them", "t0": 0, "t_end": 3, "text": "Tell me about a time you led a team", "qtype": "behavioural", "qmin": 90, "qmax": 120}
    long_text = "I think maybe it was kind of a hard project and I guess we probably did some things, you know. " * 6
    rows = [q, {"id": 2, "source": "me", "t0": 5, "t_end": 45, "text": long_text, "extra": 1}]
    rows += [{"id": 3, "source": "me", "t0": 46, "t_end": 86, "text": long_text, "extra": 1},
             {"id": 4, "source": "me", "t0": 87, "t_end": 127, "text": long_text, "extra": 1},
             {"id": 5, "source": "me", "t0": 128, "t_end": 168, "text": long_text, "extra": 1}]
    c = app.check_answer(rows, q)
    assert "long" in c["flags"] and "hedging" in c["flags"] and "no_example" in c["flags"], c
    good = [q, {"id": 2, "source": "me", "t0": 5, "t_end": 60,
                "text": "At my last company I led a team of six during a migration. We moved forty databases in three months and reduced downtime by 90 percent, "
                        "which the management noticed. I set the plan, split the work and reviewed every cutover with the team.", }]
    c = app.check_answer(good, q)
    assert c["flags"] == ["good"], c
    assert app.check_answer([q], q) is None


def test_notes_are_cleaned(app):
    n = app.clean_notes({"me_facts": ["8 years OSPF", "", 5, {"x": 1}], "topics": "not a list", "brief": {"key_messages": ["a", "b"], "junk": 1},
                         "commitments": ["send the report on Friday"] * 30})
    assert n["me_facts"] == ["8 years OSPF", "5"] and "topics" not in n
    assert len(n["commitments"]) == 8 and n["brief"] == {"key_messages": ["a", "b"]}
    assert app.clean_notes("x") == {} and app.clean_notes(None) == {}


def test_coach_writes_notes_and_alert(app):
    import time
    eng, sess, calls = _engine(app)
    now = time.time()
    sess.add("them", now - 90, now - 85, "How many years of OSPF experience do you have", "en", [], None)
    sess.add("me", now - 80, now - 70, "I have five years of OSPF experience", "en", [], None)
    sess.add("them", now - 30, now - 25, "And how long did you work with BGP", "en", [], None)
    sess.add("me", now - 20, now - 10, "Eight years of OSPF and BGP", "en", [], None)
    published = []
    eng.hub = type("H", (), {"publish": lambda self, *a, **k: published.append((a[0], k)), "toast": lambda *a, **k: None})()
    eng.chat_targets = lambda task: [("groq", "a", None), ("groq", "b", None)]
    seen = {}

    def chat(chain, messages, *a, **k):
        seen["sys"], seen["user"] = messages[0]["content"], messages[1]["content"]
        return ('{"phase": "technical", "topic": "OSPF", "difficulty": 3, "trend": "same", "signal": "neutral", '
                '"tip": "Clarify: five or eight years?", "alert": "You said 5 years, then 8 years of OSPF", '
                '"notes": {"me_facts": ["5 years OSPF (said first)", "8 years OSPF and BGP (said later)"], "topics": ["OSPF", "BGP"]}}')
    eng.run_chat = chat
    eng._coach_run()
    assert eng.coach_state["alert"].startswith("You said 5 years")
    assert sess.coach_notes["me_facts"][0].startswith("5 years") and "coach_notes" in [p[0] for p in published]
    assert "interview" in seen["sys"].lower() or "meeting" in seen["sys"].lower()
    assert "Tips already shown" not in seen["user"]
    # (second run shows the earlier tip so it is not repeated)
    eng.coach_last = 0
    eng._coach_run()
    assert "Clarify: five or eight years?" in seen["user"]
    assert "5 years OSPF" in eng.notes_text()


def test_coach_help_and_brief(app):
    import time
    eng, sess, calls = _engine(app)
    eng.cfg["about_me"] = "Network engineer, 8 years OSPF and BGP, Oracle DBA."
    eng.chat_targets = lambda task: [("groq", "a", None), ("groq", "b", None)]
    got = {}

    def chat(chain, messages, *a, **k):
        got["chain"], got["user"] = chain, messages[-1]["content"]
        if "prepare a user" in messages[0]["content"]:
            return '{"key_messages": ["Show OSPF depth"], "strengths": ["8 years"], "risks": ["No cloud: say what you learn"], "ask_them": ["What does success look like?"]}'
        return "  Say the direct answer first, then one example.  "
    eng.run_chat = chat
    eng.coach_brief()
    assert sess.coach_notes["brief"]["key_messages"] == ["Show OSPF depth"]
    assert "8 years OSPF" in got["user"]
    now = time.time()
    sess.add("them", now - 10, now - 5, "Explain OSPF areas", "en", [], None)
    out = eng.coach_help("what now?")
    assert out == "Say the direct answer first, then one example."
    assert len(got["chain"]) == 2 and "what now?" in got["user"]         # the help key uses the whole chain
    eng.stop_event.set()


def test_question_types_negatives_and_german(app):
    not_behavioural = ["Have you ever used Terraform?", "Can you give me an example of an index?", "How do you resolve a merge conflict in git?",
                       "What happens on a node failure in a cluster?", "How does RMAN handle a block failure?"]
    for t in not_behavioural:
        assert app.classify_question(t)[0] != "behavioural", t
    assert app.classify_question("Have you ever used Terraform?")[0] == "experience"
    assert app.classify_question("How much do you know about our company?")[0] != "salary"
    assert app.classify_question("How much do you expect to earn?")[0] == "salary"
    assert app.classify_question("What design patterns do you know?")[0] != "design"
    assert app.classify_question("How would you tune a slow SQL query?")[0] not in ("coding", "behavioural")
    assert app.classify_question("Have you ever had a conflict with a colleague?")[0] == "behavioural"
    assert app.classify_question("Wie gehen Sie mit Konflikten um?")[0] == "behavioural"
    assert app.classify_question("Wie viel Gehalt erwarten Sie?")[0] == "salary"
    assert app.classify_question("Stellen Sie sich bitte kurz vor")[0] == "intro"
    assert app.classify_question("Haben Sie noch Fragen an uns?")[0] == "candidate_q"


def test_reply_check_scope_and_speed(app):
    q1 = {"id": 1, "source": "them", "t0": 0, "t_end": 3, "text": "Tell me about a time you led a team", "qtype": "behavioural", "qmin": 90, "qmax": 120}
    ans = "At my last job I led six people and we cut the downtime by 90 percent in three months. " * 5
    q2 = {"id": 3, "source": "them", "t0": 200, "t_end": 204, "text": "Where are you based at the moment please"}
    late = {"id": 4, "source": "me", "t0": 206, "t_end": 209, "text": "I am in Berlin now, close to the office."}
    rows = [q1, {"id": 2, "source": "me", "t0": 5, "t_end": 65, "text": ans}, q2, late]
    c = app.check_answer(rows, q1)
    assert c["secs"] == 60 and "long" not in c["flags"], c                # the later short reply is not added to the first answer
    c2 = app.check_answer(rows, q2)
    assert c2["words"] < 12 and c2["qid"] == 3
    one = [q1, {"id": 2, "source": "me", "t0": 5, "t_end": 95, "text": "word " * 150}]
    c = app.check_answer(one, q1)
    assert c["secs"] == 90 and c["wpm"] == 100 and "fast" not in c["flags"], c   # a long single line keeps its real duration
    de = [q1, {"id": 2, "source": "me", "t0": 5, "t_end": 60, "lang": "de", "text": "Ich habe sechs Leute geführt und wir haben das Projekt pünktlich abgeschlossen. " * 3}]
    assert "no_example" not in app.check_answer(de, q1)["flags"]         # English word lists: no verdict on German


def test_notes_merge_and_stale_check(app):
    import time
    eng, sess, calls = _engine(app)
    now = time.time()
    sess.add("them", now - 90, now - 85, "How many years of OSPF experience do you have", "en", [], None)
    sess.add("them", now - 60, now - 55, "And how long did you work with BGP", "en", [], None)
    sess.coach_notes = {"me_facts": ["8 years OSPF"], "topics": ["OSPF"], "brief": {"key_messages": ["x"]}}
    eng.hub = type("H", (), {"publish": lambda self, *a, **k: None, "toast": lambda *a, **k: None})()
    eng.chat_targets = lambda task: [("groq", "a", None)]
    eng.run_chat = lambda *a, **k: '{"phase": "technical", "difficulty": 1e999, "alert": "None", "notes": {"topics": ["BGP"]}}'
    eng._coach_run()
    assert sess.coach_notes["me_facts"] == ["8 years OSPF"] and sess.coach_notes["topics"] == ["BGP"]   # an omitted list is kept
    assert sess.coach_notes["brief"]["key_messages"] == ["x"] and eng.coach_state["alert"] == ""
    eng.last_check = {"qid": 1, "secs": 10, "words": 20, "wpm": None, "flags": []}
    assert eng.current_check() is None                                    # it belongs to an older question
    eng.chat_targets = lambda task: [("llm", "local", None), ("groq", "a", None)]
    assert eng.coach_chain() == [("llm", "local", None)]                  # the background text stays on the local model
    eng.chat_targets = lambda task: []
    assert eng.coach_chain() == []


def test_help_key_is_single_flight(app):
    eng, sess, calls = _engine(app)
    class A: pass
    a = app.App.__new__(app.App)
    a.engine = eng
    a.cfg = eng.cfg                                              # (the app's settings: the Features switches live there)
    eng.cfg["coach"] = True
    assert eng._help_lock.acquire(blocking=False)
    r = a.api_coach_help("x")
    assert not r["ok"] and "still" in r["error"]
    eng._help_lock.release()
    eng.cfg["coach"] = False
    assert "off" in a.api_coach_help("x")["error"]


def test_exam_mode_and_describe_questions(app):
    assert "exam" in app.MODES and "client" in app.MODES and app.CHOICES["meeting_mode"] == tuple(app.MODES)
    assert "IELTS" in app.MODES["exam"]["name"] and "level" in app.MODES["exam"]["answer"]
    assert app.classify_question("Describe a place you like to visit")[0] == "describe"
    assert app.classify_question("You should say where it is and why you like it")[0] == "describe"
    assert app.classify_question("Describe a time you failed at work")[0] == "behavioural"
    for t in ("Now I'd like you to talk about your hometown.", "Do you agree or disagree with this statement", "Let's move to part two"):
        assert app.looks_askable(t, "en"), t
    assert "speaking-test" in app.COACH_ROLES["exam"] and "negotiation" in app.COACH_ROLES["client"]


def test_overlay_api_off_windows(app):
    a = app.App.__new__(app.App)
    a.cfg = dict(app.DEFAULTS)
    a.hub = type("H", (), {"publish": lambda self, *x, **k: None})()
    a.url = "http://127.0.0.1:1/?t=x"
    a.closing = __import__("threading").Event()
    a.overlay = app.Overlay(a)
    real = app.sys.platform
    app.sys.platform = "linux"                                  # never open a real window from a test, also on Windows
    try:
        r = a.api_overlay()
        assert r["ok"] is False and "Windows" in r["error"] and r["on"] is False
        assert a.api_overlay(False)["ok"] is True
    finally:
        app.sys.platform = real
    assert app.DEFAULTS["overlay_alpha"] == 65 and app.DEFAULTS["overlay_hide"] is True and app.CHOICES["overlay_pos"] == ("top", "bottom")


def _chat_returns(eng, text, seen=None):
    eng.chat_targets = lambda task: [("groq", "a", None)]
    def chat(chain, messages, *a, **k):
        if seen is not None:
            seen.append(messages)
        return text
    eng.run_chat = chat


def test_scores_clean_and_saved(app, tmp_path):
    sc = app.clean_scores({"clarity": 4, "structure": "3", "depth": 9, "signals": 0, "follow_up": 2.6, "fixes": ["a", "", 5, {"x": 1}, "b", "c", "d"], "best": " ok "})
    assert sc["clarity"] == 4 and sc["structure"] == 3 and sc["depth"] == 5 and "signals" not in sc and sc["follow_up"] == 3
    assert sc["fixes"] == ["a", "5", "b"] and sc["best"] == "ok" and sc["overall"] == round((4 + 3 + 5 + 3) / 4, 1)
    assert app.clean_scores({"clarity": 4}) == {} and app.clean_scores("x") == {} and app.clean_scores({"clarity": float("inf"), "structure": 1, "depth": 1}) == {}
    eng, sess, calls = _engine(app)
    import time
    now = time.time()
    for i, (src, t) in enumerate([("them", "Explain OSPF areas to me please"), ("me", "OSPF areas split the network so that routers keep small tables"),
                                  ("them", "And what is a stub area exactly"), ("me", "A stub area blocks external routes and uses a default route")]):
        sess.add(src, now - 100 + i * 10, now - 95 + i * 10, t, "en", [], None)
    _chat_returns(eng, 'Here: {"clarity": 4, "structure": 3, "depth": 4, "signals": 2, "follow_up": 3, "fixes": ["Add a number"], "best": "OSPF answer"}')
    got = eng.debrief_scores(sess.ordered(), {})
    assert got["overall"] == 3.2 and got["fixes"] == ["Add a number"]
    sess.set_scores(got)
    assert sess.state()["scores"]["clarity"] == 4
    sess.set_feedback("new review")
    assert sess.scores == {}                                   # the old scores do not belong to a new review


def test_job_prep_and_note(app):
    eng, sess, calls = _engine(app)
    eng.cfg["job_ad"] = "Senior Oracle DBA. RMAN, Data Guard, Linux, 5+ years. " * 3
    assert "Senior Oracle DBA" in app.job_note(eng.cfg) and app.job_note({"job_ad": ""}) == ""
    assert "Senior Oracle DBA" in eng.coach_background(1000)
    seen = []
    _chat_returns(eng, "## What they look for\n- RMAN", seen)
    out = eng.job_prep(eng.cfg["job_ad"])
    assert out.startswith("## What they look for") and "Job advert" in seen[0][1]["content"]
    a = app.App.__new__(app.App)
    a.engine = None
    a.cfg = {"job_ad": ""}
    a.tool_eng = None
    assert a.api_job_prep("short")["ok"] is False


def test_practice_questions_grade_and_finish(app, tmp_path):
    import time
    app.PRACTICE_PATH = str(__import__("pathlib").Path(tmp_path) / "practice.json")
    eng, sess, calls = _engine(app)
    _chat_returns(eng, 'Sure: ["Tell me about yourself", "Why this job?", "Tell me about yourself", 5, ""]')
    qs = eng.practice_questions("general", 5)
    assert [q["q"] for q in qs] == ["Tell me about yourself", "Why this job?"] and qs[0]["type"] == "intro"
    _chat_returns(eng, "1. What is RMAN?\n2. Explain Data Guard\nnot a question")
    assert [q["q"] for q in eng.practice_questions("technical", 3)] == ["What is RMAN?", "Explain Data Guard"]
    app.save_practice({"weak": [{"q": "Old weak question?", "t": 1, "type": ""}], "history": []})
    _chat_returns(eng, '["New one?"]')
    q = eng.practice_questions("redrill", 3)
    assert q[0]["q"] == "Old weak question?" and q[1]["q"] == "New one?"
    # grading
    good = ("At my last company I migrated forty databases to a new RMAN scheme in three months and cut the restore time by half. " * 2)
    _chat_returns(eng, '{"clarity": 4, "structure": 4, "depth": 3, "signals": 5, "tips": ["Shorter"], "better": "Say it in one line.", "follow_up": "How did you test it?"}')
    r = eng.practice_grade("Tell me about a time you improved backups", good, "behavioural", 45)
    assert r["avg"] == 4.0 and r["follow_up"].startswith("How") and r["scores"]["signals"] == 5 and not r["empty"]
    assert eng.practice_grade("Why this job?", "um", "general", 3)["empty"] is True      # no model call for no answer
    _chat_returns(eng, "no json here")
    try:
        eng.practice_grade("Why this job?", good, "general", 40)
        assert False
    except app.APIError:
        pass
    _chat_returns(eng, '{"clarity": 4, "structure": 4, "depth": 3, "signals": 5, "language": 0}')
    r = eng.practice_grade("Describe a place", good, "english", 40)    # one missing score gets the average of the others
    assert r["scores"]["language"] == 4
    _chat_returns(eng, '{"clarity": 4, "structure": 0, "depth": 0, "signals": 5}')
    try:
        eng.practice_grade("Describe a place", good, "english", 40)     # too many missing scores are not accepted
        assert False
    except app.APIError:
        pass
    # finish: good answers leave the weak list, bad ones and empty ones enter it
    a = app.App.__new__(app.App)
    a.practice_until = time.time() + 100
    out = a.api_practice_finish("general", [{"q": "Old weak question?", "avg": 4.2}, {"q": "Bad one?", "avg": 2.0}, {"q": "Skipped?", "avg": None},
                                            {"q": 5}, "x", {"q": "Bool?", "avg": True}])
    st = app.load_practice()
    assert out["avg"] == 2.7 or out["avg"] == 3.1
    assert {w["q"] for w in st["weak"]} == {"Bad one?", "Skipped?", "Bool?"} and a.practice_until == 0.0
    assert st["history"][-1]["kind"] == "general"


def test_practice_mutes_answers_and_coach(app):
    import time
    eng, sess, calls = _engine(app)
    eng.app.practice_until = time.time() + 60
    try:
        _say(app, eng, "Tell me about a time you led a team through an outage?")
        time.sleep(0.6)
        assert calls == [] and eng.plan is None
        eng.app.practice_until = 0.0
        _say(app, eng, "Tell me about a time you led a team through another outage?")
        time.sleep(0.8)
        assert len(calls) == 1
    finally:
        eng.stop_event.set()


def test_overlay_keys_and_new_defaults(app):
    for k in ("toggle", "more", "less", "screen", "answer", "coach", "up", "down", "prev", "next"):
        assert k in app.OV_KEYS
    assert app.DEFAULTS["auto_screen"] is False and app.DEFAULTS["debrief_scores"] is True and app.DEFAULTS["job_ad"] == ""


def test_second_piece_after_the_first_was_answered(app):
    import time
    eng, sess, calls = _engine(app)
    try:
        _say(app, eng, "Tell me about a time")
        time.sleep(1.5)                                          # the first piece is answered already (no plan is waiting)
        assert len(calls) == 1
        _say(app, eng, "when you led a team through an outage")  # used to raise TypeError (old plan is None)
        time.sleep(1.8)
        assert len(calls) == 2 and "outage" in calls[1][2] and "Tell me about a time" in calls[1][2]
    finally:
        eng.stop_event.set()


def test_json_from_and_clip_and_practice_file(app, tmp_path):
    assert app.json_from('Scores {see} {"a": 1}') == {"a": 1}
    assert app.json_from('[1] refs ["q?"]', list) == ["q?"]                # a stray [1] is skipped
    assert app.json_from("no json") is None and app.json_from(None) is None and app.json_from('{"a": ', dict) is None
    assert app.clip("", 10) == "" and app.clip(None, 10) == "" and app.clip("a  b", 10) == "a b" and app.clip("x" * 20, 5) == "xxxxx…"
    sc = app.clean_scores({"clarity": 4, "structure": 4, "depth": 4})
    assert sc["best"] == "" and sc["fixes"] == []                          # a missing text stays empty (never the word "str")
    assert app.score_int(True) == 0 and app.score_int(0) == 0 and app.score_int(8) == 5 and app.score_int("x") == 0
    path = __import__("pathlib").Path(tmp_path) / "practice.json"
    app.PRACTICE_PATH = str(path)
    path.write_text('{"weak": [{"q": "Ok?", "t": "bad"}, {"q": 5}, "x"], "history": [{"t": "x", "avg": "y", "n": "z"}, 5]}', encoding="utf-8")
    st = app.load_practice()
    assert st["weak"] == [{"q": "Ok?", "t": 0.0, "type": ""}] and st["history"][0]["avg"] is None
    a = app.App.__new__(app.App)
    a.practice_until = 0.0
    out = a.api_practice_finish("general", [{"q": "N?", "avg": float("nan")}, {"q": "I?", "avg": float("inf")}, {"q": "Fine?", "avg": 9}])
    assert out["ok"] and out["avg"] == 5.0                                 # only usable numbers count, and never above 5
    path.write_text("not json", encoding="utf-8")
    assert app.load_practice() == {"weak": [], "history": []}


def test_practice_mute_leaves_no_thinking_line(app):
    import time
    eng, sess, calls = _engine(app)
    eng.app.practice_until = time.time() + 60
    try:
        _say(app, eng, "Tell me about a time you led a team through an outage?")
        time.sleep(0.5)
        row = [r for r in sess.ordered() if r["source"] == "them"][0]
        assert row.get("ans_state") == "none" and calls == []
    finally:
        eng.stop_event.set()


def test_job_ad_is_capped(app):
    assert len(app.sanitize({"job_ad": "x" * 20000})["job_ad"]) == 8000


def test_round2_fixes(app):
    assert app.score_int("4/5") == 4 and app.score_int(" 3 / 5") == 3
    assert app.reply_numbers("hello there", float("inf"))["secs"] == 0
    a = app.App() if hasattr(app, "App") else None
    if a is not None:
        assert a.api_practice_questions(kind=["x"], n=float("inf"))["ok"] is False
        assert a.api_practice_grade(question="Q?", answer="a", kind={"a": 1}, secs=float("inf")).get("ok") in (True, False)
        a.running, a.started_at = True, __import__("time").time()
        assert a.api_practice_mode(on=True).get("ignored") is True and a.practice_until == 0.0


def test_selfcheck_and_new_fixes(app):
    a = app.App()
    r = a.api_selfcheck()
    assert r["ok"] and any(i["name"] == "Program" for i in r["items"]) and "Program" in r["report"]
    assert all(i["state"] in ("ok", "warn", "bad", "info") for i in r["items"])
    assert a.api_selfcheck(overlay_live=True)["ok"]                          # off Windows: skipped, no window
    assert app.clean_key(" “gsk_abc‏123” \n") == "gsk_abc123" and app.clean_key("گ") == ""
    assert app.sanitize({"api_key": "gsk_‏x y"})["api_key"] == "gsk_xy"


def test_blank_reply_tries_next_model_and_preview_is_soft(app):
    eng, sess, calls = _engine(app)
    eng.app.get_api = lambda pid=None: object()
    eng.app.note_request = lambda: None
    outs = iter(["", "second model text"])
    orig = app.chat_compat
    app.chat_compat = lambda *a, **k: next(outs)
    try:
        chain = [("groq", "m1", None), ("groq", "m2", None)]
        assert eng.run_chat(chain, [], 10, 0.1, lambda t: None) == "second model text"
        eng.models.cooldown("groq|m1", 60)
        try:
            eng.run_chat(chain, [], 10, 0.1, lambda t: None, soft=True)      # a preview does not use a resting model
            assert False
        except app.APIError:
            pass
        assert not eng.models.ok("groq|m1") and eng.models.ok("groq|m2")
    finally:
        app.chat_compat = orig
        eng.stop_event.set()


def test_session_save_survives_odd_text_and_tells(app, tmp_path):
    told = []
    app.save_alert = lambda kind, msg: told.append((kind, msg))
    eng, sess, calls = _engine(app)
    sess.add("them", 1.0, 2.0, "hello \ud83d there", "en", [1])
    sess.dirty = True
    sess._save()                                                              # a lone surrogate does not lose the file
    assert os.path.isfile(sess.path)
    eng.stop_event.set()


def test_coach_gives_the_words_to_say(app):
    eng, sess, calls = _engine(app)
    eng.chat_targets = lambda task: [("groq", "a", None)]
    seen = {}

    def chat(chain, messages, *a, **k):
        seen["system"] = messages[0]["content"]
        return ('{"phase": "intro", "topic": "x", "difficulty": 2, "trend": "same", "signal": "neutral", '
                '"tip": "بگو رشته‌ات", "say": "I graduated in computer engineering.", "alert": "", "notes": {}}')
    eng.run_chat = chat
    now = __import__("time").time()
    sess.add("them", now - 20, now - 15, "Where did you study?", "en", [], None)
    sess.add("them", now - 10, now - 5, "And your degree?", "en", [], None)
    eng._coach_run()
    assert "English" in seen["system"]       # the words to say are asked for in the meeting language
    assert eng.coach_state["say"] == "I graduated in computer engineering."
    eng.coach_help("what now?")
    assert "SAY:" in seen["system"] and "English" in seen["system"]
    eng.stop_event.set()


def test_coach_situation_sees_silence_repeats_and_kinds(app):
    import time
    eng, sess, calls = _engine(app)
    now = time.time()
    eng.cfg["transcribe_me"] = True
    eng.app.captures = [type("Cap", (), {"source": "me"})()]
    e = sess.add("them", now - 14, now - 10, "What are your salary expectations for this role?", "en", [], None)
    sit, stuck, _q = eng.coach_situation()
    assert stuck and "NOT started answering" in sit and "Pay question" in sit
    sess.add("me", now - 5, now - 1, "I am looking for a range of", "en", [], None)
    sit, stuck, _q = eng.coach_situation()
    assert not stuck and "answered for about" in sit
    assert app.coach_intents("What is the cache hit rate and the single point of failure in high availability?") == []
    eng.cfg["meeting_mode"] = "lecture"
    sess.add("them", now - 9, now - 8, "Why does the optimizer pick this plan?", "en", [], None)
    assert eng.coach_situation()[1] is False                # a class: nobody expects the user to answer
    assert [n for n, _ in app.coach_intents("Do you have any questions for us?")] == ["your_questions"]
    assert app.coach_intents("We use Oracle here.") == []
    eng.stop_event.set()


def test_coach_rescues_a_stuck_user_once(app):
    import time
    eng, sess, calls = _engine(app)
    eng.cfg["coach"] = True
    eng.cfg["transcribe_me"] = True
    eng.app.captures = [type("Cap", (), {"source": "me"})()]
    runs = []
    eng._coach_run = lambda: runs.append(time.time())
    now = time.time()
    sess.add("them", now - 12, now - 9, "Can you explain how Data Guard switchover works?", "en", [], None)
    eng.coach_dirty = False
    waits = iter([False, False, True])
    eng.stop_event.wait = lambda t: next(waits)
    eng.coach_last = now - 20
    eng._coach_loop()
    assert len(runs) == 1                                   # silent after a question: one look at once, not twice
    eng.me_voice_at = time.time()                            # the mic hears the user: not stuck
    assert eng.coach_situation()[1] is False
    eng.stop_event.set()


def _news_fakes():
    md = {"openai": {"name": "OpenAI", "doc": "https://platform.openai.com/docs", "models": {
              "gpt-9": {"name": "GPT-9", "release_date": "2026-09-20", "modalities": {"input": ["text", "image"], "output": ["text"]},
                        "cost": {"input": 2.5, "output": 10}, "limit": {"context": 400000}},
              "gpt-9-transcribe": {"name": "GPT-9 Transcribe", "release_date": "2026-09-25", "modalities": {"input": ["audio"], "output": ["text"]},
                                   "cost": {"input": 3, "output": 6}},
              "img-x": {"name": "Image X", "release_date": "2026-09-25", "modalities": {"input": ["text"], "output": ["image"]}}}},
          "azure": {"name": "Azure", "models": {"gpt-9": {"name": "GPT-9", "release_date": "2026-09-21",
                    "modalities": {"input": ["text"], "output": ["text"]}, "cost": {"input": 2.5, "output": 10}}}},
          "newco": {"name": "NewCo", "api": "https://api.newco.ai/v1", "models": {"nc-1": {"name": "NC-1", "release_date": "2026-09-28",
                    "modalities": {"input": ["text"], "output": ["text"]}, "cost": {"input": 0, "output": 0}, "open_weights": True}}},
          "broken": "x"}
    orr = {"data": [{"id": "openai/gpt-9", "name": "OpenAI: GPT-9", "created": 1790300000, "context_length": 400000,
                     "architecture": {"input_modalities": ["text"], "output_modalities": ["text"]},
                     "pricing": {"prompt": "0.0000025", "completion": "0.00001"}},
                    {"id": "meta/llama-x:free", "name": "Meta: Llama X (free)", "created": 1790400000,
                     "architecture": {"input_modalities": ["text"], "output_modalities": ["text"]}, "pricing": {"prompt": "0", "completion": "0"}},
                    "junk"]}
    hf = [{"id": "someone/whisper-x-ct2", "createdAt": "2026-09-27T10:00:00.000Z", "tags": ["ctranslate2"], "likes": 5},
          {"id": "lab/new-asr", "createdAt": "2026-09-26T10:00:00.000Z", "tags": ["transformers"]}, {"id": "../bad"}]
    return md, orr, hf


def test_model_news_parse_group_and_use(app):
    import datetime
    md, orr, hf = _news_fakes()
    conn = {"openai": "openai2"}
    src = app.parse_models_dev(md, conn) + app.parse_openrouter(orr, conn) + app.parse_hf_asr(hf)
    assert not any(x["model"] == "img-x" for x in src)                        # pictures are not for this program
    g = {x["key"]: x for x in app.group_news(src, 365, datetime.date(2026, 10, 3))}
    gpt = g["gpt-9"]
    assert gpt["kind"] == "chat" and gpt["day"] == "2026-09-20" and len(gpt["sources"]) == 3
    assert gpt["sources"][0]["connected"] == "openai2"                        # the service you have a key for comes first
    assert next(s for s in gpt["sources"] if s["provider"] == "azure")["usable"] == "no"
    orr_src = next(s for s in gpt["sources"] if s["provider"] == "openrouter")
    assert orr_src["in"] == 2.5 and orr_src["out"] == 10.0
    assert g["gpt-9-transcribe"]["kind"] == "stt"
    assert g["llama-x"]["free"] and g["nc-1"]["free"] and g["nc-1"]["sources"][0]["usable"] == "custom"
    assert g["nc-1"]["sources"][0]["base_url"] == "https://api.newco.ai/v1"
    assert g["hf:someone/whisper-x-ct2"]["sources"][0]["usable"] == "local" and g["hf:lab/new-asr"]["sources"][0]["usable"] == "no"
    old = app.group_news(src, 7, datetime.date(2027, 6, 1))
    assert old == []                                                            # older than the window: not listed


def test_model_news_api_marks_new_and_survives_offline(app, tmp_path):
    import os
    md, orr, hf = _news_fakes()
    app.MODEL_NEWS_PATH = os.path.join(str(tmp_path), "news.json")
    a = app.App()
    calls = []

    def fake(proxy, connected, get=None):
        calls.append(1)
        return (app.parse_models_dev(md, connected) + app.parse_openrouter(orr, connected), [])
    orig = app.fetch_model_news
    app.fetch_model_news = fake
    try:
        r = a.api_model_news(refresh=True, days=3650)
        assert r["ok"] and r["new"] == 0 and r["items"]                         # the first look: nothing is 'new'
        md["openai"]["models"]["gpt-10"] = {"name": "GPT-10", "release_date": "2026-09-30",
                                            "modalities": {"input": ["text"], "output": ["text"]}}
        r = a.api_model_news(refresh=True, days=3650)
        assert r["new"] == 1 and next(g for g in r["items"] if g["key"] == "gpt-10")["new"]
        r = a.api_model_news(seen=True, days=3650)
        r = a.api_model_news(days=3650)
        assert r["new"] == 0 and len(calls) == 2                                 # kept for 12 hours: no new download
        app.fetch_model_news = lambda *x, **k: (_ for _ in ()).throw(app.APIError("offline"))
        r = a.api_model_news(refresh=True, days=3650)
        assert r["ok"] and r["items"] and "older list" in r["failed"][0]         # offline: the last list stays usable
    finally:
        app.fetch_model_news = orig


def test_model_news_links_are_safe(app):
    bad = {"evil": {"name": "Evil", "api": "javascript:alert(1)", "doc": "javascript:alert(2)", "models": {
        "m": {"name": "M", "release_date": "2026-09-30", "modalities": {"input": ["text"], "output": ["text"]}}}},
           "odd": {"name": "Odd", "api": "http://insecure.example/v1", "doc": "file:///c:/x", "models": {
        "m2": {"name": "M2", "release_date": "2026-09-30", "modalities": {"input": ["text"], "output": ["text"]}}}}}
    src = app.parse_models_dev(bad, {})
    assert all(x["site"] == "" and x["base_url"] == "" and x["usable"] == "no" for x in src)
    assert app._https("https://api.x.ai/v1") == "https://api.x.ai/v1" and app._https("https://a.b/\"><script>") == ""


def test_model_news_odd_items_and_custom_match(app, tmp_path):
    import os
    odd = {"p": {"name": "P", "api": "https://api.p.ai/v1", "models": {
        "a": {"name": "A", "release_date": "2026-09-30", "modalities": {"input": 5, "output": ["text"]}},
        "b": {"name": "B", "release_date": "2026-09-30", "modalities": {"input": ["text"], "output": ["text"]}}}}}
    src = app.parse_models_dev(odd, {})
    assert [x["model"] for x in src] == ["b"]                                 # one odd entry does not lose the list
    assert app.parse_hf_asr([{"id": "a/b", "tags": 7, "likes": "9" * 5000}])[0]["likes"] == 0
    assert app.parse_openrouter({"data": [{"id": "x/y", "architecture": {"input_modalities": 3}}]}, {}) == []
    app.MODEL_NEWS_PATH = os.path.join(str(tmp_path), "news.json")
    a = app.App()
    a.cfg["providers"] = [{"id": "aval", "preset": "custom", "name": "AvalAI", "base_url": "https://api.avalai.ir/v1", "api_key": "k", "models": []}]
    orig = app.fetch_model_news
    app.fetch_model_news = lambda proxy, conn, get=None: (app.parse_models_dev(odd, conn), [])
    try:
        r = a.api_model_news(refresh=True, days=3650)
        assert r["ok"] and all(not s["connected"] for g in r["items"] for s in g["sources"])   # another address: not 'your key'
        a.cfg["providers"][0]["base_url"] = "https://api.p.ai/v1"
        r = a.api_model_news(days=3650)
        assert r["items"][0]["sources"][0]["connected"] == "aval"
    finally:
        app.fetch_model_news = orig


def test_vtt_export_and_keep_awake(app):
    import types
    rows = [{"id": 1, "source": "them", "t0": 1_700_000_100.0, "t_end": 1_700_000_103.5, "text": "Hello there", "translation": "سلام", "speaker": None}]
    s = types.SimpleNamespace(ordered=lambda: rows, started=__import__("datetime").datetime.fromtimestamp(1_700_000_099.0), recording=False,
                              label=lambda r: "Them")
    v = app.export_bytes(s, "vtt").decode("utf-8")
    assert v.startswith("WEBVTT\n\n") and "00:00:01.000 --> 00:00:04.500" in v and "," not in v.split("\n")[2]
    assert "Hello there" in v and "Them:" not in v and "سلام" in v          # one speaker: no name on every cue
    app.keep_awake(True); app.keep_awake(False)                                # off Windows: does nothing, never fails


def test_subtitles_are_short_and_timed(app):
    long = ("First I would confirm that the primary database is really down and not just a network problem, because a "
            "false failover is worse than a short wait. Then I would fail over to the standby with the broker, check that "
            "the apply lag was zero, and point the application to the new primary.")
    rows = [{"source": "them", "t0": 10.0, "t_end": 40.0, "text": long, "translation": "اول مطمئن می‌شوم " * 12, "speaker": None},
            {"source": "me", "t0": 41.0, "t_end": 43.0, "text": "Okay.", "translation": "", "speaker": None}]
    srt = app.srt_text(rows, 0.0, lambda r: "Them" if r["source"] == "them" else "Me")
    cues = [c for c in srt.strip().split("\n\n")]
    assert len(cues) >= 4
    times = []
    for c in cues:
        ln = c.split("\n")
        a, b = ln[1].split(" --> ")
        sec = lambda x: int(x[:2]) * 3600 + int(x[3:5]) * 60 + float(x[6:].replace(",", "."))
        times.append((sec(a), sec(b)))
        assert all(len(t) <= 42 or " " not in t for t in ln[2:4])            # lines of at most 42 characters
        assert sec(b) - sec(a) <= 7.5
    assert times[0][0] == 10.0 and abs(times[-2][1] - 40.0) < 0.01 and times[-1] == (41.0, 43.0)
    assert all(times[i][1] <= times[i + 1][0] + 0.001 for i in range(len(times) - 1))   # never overlapping
    assert cues[0].split("\n")[2].startswith("Them: ") and "Them:" not in cues[1]
    assert app._sub_pieces("", 84) == [] and app._sub_pieces("a b", 84) == ["a b"]


def test_whisper_cpp_engine_with_a_fake_server(app, tmp_path):
    import os, stat, sys, time
    if sys.platform == "win32":
        return                                                    # the fake server below is a shell script
    srv = os.path.join(str(tmp_path), "whisper-server")
    with open(srv, "w") as f:
        f.write("#!" + sys.executable + "\n" + r'''
import sys, json, http.server
a = sys.argv; port = int(a[a.index("--port") + 1])
print("ggml_vulkan: 0 = Intel(R) Iris(R) Xe Graphics (Intel Corporation) | uma: 1", flush=True)
class H(http.server.BaseHTTPRequestHandler):
    def log_message(self, *x): pass
    def do_GET(self):
        self.send_response(200); self.end_headers(); self.wfile.write(b'{"status":"ok"}')
    def do_POST(self):
        n = int(self.headers["Content-Length"]); body = self.rfile.read(n)
        lang = b'name="language"\r\n\r\nen' in body
        out = {"language": "english", "text": " Hello world.", "segments": [{"text": " Hello world.", "start": 0.0, "end": 1.2,
               "avg_logprob": -0.2, "no_speech_prob": 0.01}]} if lang else {"error": "no language"}
        self.send_response(200); self.end_headers(); self.wfile.write(json.dumps(out).encode())
http.server.HTTPServer(("127.0.0.1", port), H).serve_forever()
''')
    os.chmod(srv, os.stat(srv).st_mode | stat.S_IEXEC)
    model = os.path.join(str(tmp_path), "ggml-small-q5_1.bin")
    with open(model, "wb") as f:
        f.write(b"lmgg" + b"\0" * 100)
    old = os.environ.get("MA_WHISPER_SERVER")
    os.environ["MA_WHISPER_SERVER"] = srv
    old_data, app.DATA_DIR = app.DATA_DIR, str(tmp_path)          # put back at the end: other tests share the module
    try:
        assert app.model_folder(model) == (os.path.abspath(model), "") and app.model_label(model) == "small-q5_1 (whisper.cpp)"
        L = app.LocalWhisper()
        L.load(model, "auto")
        assert L.wait(30) and L.state == "ready", L.error
        assert "Iris" in L.device and "Vulkan" in L.device
        r = L.transcribe(app.np.zeros(16000, app.np.float32), language="en")
        assert r["text"] == "Hello world." and r["segments"][0]["end"] == 1.2 and app.clean_transcript(r) == "Hello world."
        try:
            L.transcribe(app.np.zeros(16000, app.np.float32), language=None)
            assert False
        except app.Transient as e:
            assert "no language" in str(e)
        old_pid = L.model.proc.pid
        L.model.proc.kill()                                          # the engine dies (driver reset, out of memory)
        L.model.proc.wait(5)
        r = L.transcribe(app.np.zeros(16000, app.np.float32), language="en")
        assert r["text"] == "Hello world." and L.model.proc.pid != old_pid    # started again by itself
        pid = L.model.proc.pid
        L.unload()
        time.sleep(1.5)
        assert not os.path.exists(f"/proc/{pid}") or open(f"/proc/{pid}/stat").read().split()[2] == "Z"   # the engine ends
    finally:
        app.DATA_DIR = old_data
        if old is None:
            os.environ.pop("MA_WHISPER_SERVER", None)
        else:
            os.environ["MA_WHISPER_SERVER"] = old
    os.environ.pop("MA_WHISPER_SERVER", None)
    with open(model, "wb") as f:
        f.write(b"lmgg")
    assert app.model_folder(model)[0] == "" and "whisper.cpp engine is not in this copy" in app.model_folder(model)[1]


def test_fast_file_mode_with_the_local_model(app, tmp_path):
    import os, time
    np = app.np
    sr = 16000
    calls = []

    def blocks():
        for k in range(72):                                     # 12 minutes in 10 s blocks: speech, then 1 s of silence
            b = (np.random.default_rng(k).standard_normal(10 * sr) * 0.1).astype(np.float32)
            b[-sr:] = 0.0
            yield b

    hiccup = [1]

    def long(audio, language=None, prompt=None):
        if len(calls) == 1 and hiccup:                          # one short problem on the second piece: tried again
            hiccup.pop()
            raise app.Transient("Local model: connection reset")
        calls.append((len(audio) / sr, language, prompt))
        n = len(audio) / sr
        return ([{"start": 1.0, "end": 4.0, "text": f" Part {len(calls)} starts here.", "no_speech_prob": 0.01, "avg_logprob": -0.2,
                  "compression_ratio": 1.1},
                 {"start": 5.0, "end": 6.0, "text": " Thank you.", "no_speech_prob": 0.9, "avg_logprob": -1.0, "compression_ratio": 1.0},
                 {"start": n - 3, "end": n - 1, "text": " End of the stretch.", "no_speech_prob": 0.02, "avg_logprob": -0.3,
                  "compression_ratio": 1.0}], "en")
    L = app.LOCAL
    saved = (L.ready, L.long_ok, L.transcribe_long, L.state, app.open_recording)
    a = app.App()
    a.cfg.update({"stt_provider": "local", "languages": ["en", "de"], "my_language": "fa", "api_key": ""})
    L.ready, L.long_ok, L.transcribe_long, L.state = (lambda: True), (lambda: True), long, "ready"
    app.open_recording = lambda path: (blocks(), 720.0)
    try:
        path = os.path.join(str(tmp_path), "talk.mp3")
        open(path, "wb").close()
        job = app.RecordingJob(a, path)
        job.engine._translate_one = lambda *x, **k: None
        job.run()
        rows = job.session.ordered()
        assert [w[0] for w in calls][:2] and all(250 <= w[0] <= 300 for w in calls[:-1]) and sum(w[0] for w in calls) == 720
        assert calls[0][1] is None and calls[1][1] == "en"                # the language is found once, then kept
        texts = [r["text"] for r in rows]
        assert "Thank you." not in " ".join(texts) and texts[0] == f"Part 1 starts here."
        assert len(rows) == 2 * len(calls)
        t0s = [r["t0"] for r in rows]
        assert t0s == sorted(t0s)
        base = rows[0]["t0"] - 1.0
        assert abs((rows[2]["t0"] - base) - (calls[0][0] + 1.0)) < 0.01    # times are positions in the recording
        assert job.state["state"] == "done" and not hiccup
    finally:
        L.ready, L.long_ok, L.transcribe_long, L.state, app.open_recording = saved
    q = np.ones(20 * sr, np.float32)
    q[int(15.5 * sr):int(15.7 * sr)] = 0
    assert abs(app.quiet_cut(q) / sr - 15.55) < 0.1


def test_voice_prints_cluster_into_people(app):
    np = app.np
    rng = np.random.default_rng(1)
    centers = [rng.standard_normal(64) for _ in range(3)]
    embs, truth = [], []
    for i in range(30):
        who = [0, 1, 0, 2, 1][i % 5]
        v = centers[who] + 0.35 * rng.standard_normal(64)
        embs.append(v / np.linalg.norm(v))
        truth.append(who)
    auto = app.cluster_voices(embs)
    assert len(set(auto)) == 3
    mapping = {}
    for g, t in zip(auto, truth):
        mapping.setdefault(g, t)
        assert mapping[g] == t                                   # every line goes to the right person
    assert auto[0] == 0 and auto[1] == 1 and auto[3] == 2        # numbered by who speaks first
    assert len(set(app.cluster_voices(embs, k=2))) == 2 and app.cluster_voices([]) == []
    assert app.cluster_voices(embs[:1]) == [0]


def test_fbank_and_speakers_for_rows(app):
    np = app.np
    sr = 16000
    t = np.arange(sr) / sr
    f = app.fbank((0.3 * np.sin(2 * np.pi * 1000 * t)).astype(np.float32))
    assert f.shape == (98, 80) and abs(float(f.mean())) < 1e-3
    peak = int(np.argmax(f.mean(axis=0) - f.mean()))               # a 1 kHz tone lands in a middle band
    assert 20 < peak < 45 and app.fbank(np.zeros(100, np.float32)).shape == (0, 80)
    # two voices: 'A' parts are loud, 'B' parts quiet; the fake print model tells them apart by loudness
    plan = [("A", 0, 4), ("B", 5, 9), ("A", 10, 14), ("B", 15, 15.5), ("B", 16, 21)]
    audio = np.zeros(22 * sr, np.float32)
    for who, a, b in plan:
        audio[int(a * sr):int(b * sr)] = (0.5 if who == "A" else 0.05) * np.sin(2 * np.pi * 220 * np.arange(int((b - a) * sr)) / sr)
    base = 1000.0
    rows = [{"id": i + 1, "t0": base + a, "t_end": base + b} for i, (_w, a, b) in enumerate(plan)]
    blocks = (audio[i:i + 7 * sr] for i in range(0, len(audio), 7 * sr))

    def prints(seg):
        loud = float(np.sqrt((seg ** 2).mean()))
        v = np.array([1.0, 0.0]) if loud > 0.1 else np.array([0.0, 1.0])
        return v
    got = app.speakers_for_rows(rows, blocks, base, prints)
    assert got == {1: 0, 2: 1, 3: 0, 4: 1, 5: 1}                   # the 0.5 s line takes the person nearest in time


def test_link_is_downloaded_then_transcribed(app, tmp_path):
    import importlib.util, os, threading, http.server, functools, time, wave
    if importlib.util.find_spec("yt_dlp") is None or not app.ffmpeg_path():
        return                                                    # (yt-dlp and ffmpeg are in the Windows release)
    np = app.np
    folder = str(tmp_path)
    with wave.open(os.path.join(folder, "talk.wav"), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000)
        w.writeframes((np.sin(np.arange(32000) / 5) * 8000).astype(np.int16).tobytes())
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(http.server.SimpleHTTPRequestHandler, directory=folder))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    old_dl, app.DOWNLOADS_DIR = app.DOWNLOADS_DIR, os.path.join(folder, "downloads")
    a = app.App()
    a.missing_tasks = lambda *x: []
    started = []
    a.api_recording = lambda **k: started.append(k) or setattr(a, "recording", None) or {"ok": True}   # the file job takes over
    try:
        assert a.api_recording_link(url="not a link")["ok"] is False
        r = a.api_recording_link(url=f"http://127.0.0.1:{srv.server_address[1]}/talk.wav", speakers=2)
        assert r["ok"], r
        for _ in range(100):
            if started:
                break
            time.sleep(0.1)
        assert started and started[0]["speakers"] == 2 and os.path.isfile(started[0]["path"])
        assert started[0]["path"].startswith(app.DOWNLOADS_DIR) and a.recording is None and started[0]["_from"] is not None
    finally:
        app.DOWNLOADS_DIR = old_dl
        srv.shutdown()


def test_meeting_sound_is_kept_and_a_line_can_be_replayed(app, tmp_path):
    import os, threading, time, httpx, types
    np = app.np
    path = os.path.join(str(tmp_path), "meeting_x.md")
    rec = app.MeetingAudio(path)
    t = time.time()
    rec.write("them", np.full(48000, 0.1, np.float32), 48000)          # 1 s at 48 kHz
    time.sleep(1.2)                                                     # the device was away for a moment
    rec.write("them", np.full(16000, 0.2, np.float32), 16000)
    rec.write("me", np.full(8000, 0.05, np.float32), 16000)
    rec.close()
    rec.write("them", np.zeros(100, np.float32), 16000)                 # after close: ignored, no error
    parts = app.load_audio_index(os.path.join(str(tmp_path), "meeting_x.audio.json"))
    them = next(p for p in parts if p["source"] == "them")
    assert 2.0 <= them["dur"] <= 2.6 and os.path.isfile(os.path.join(str(tmp_path), them["file"]))   # the gap was filled
    s = types.SimpleNamespace(path=path, recording="")
    e = {"source": "them", "t0": them["t0"] + 1.5, "t_end": them["t0"] + 2.0}
    fp, a, b = app.audio_for_line(s, e)
    assert fp.endswith(them["file"]) and abs(a - 1.2) < 0.01 and b <= them["dur"]
    assert app.audio_for_line(s, {"source": "me", "t0": t + 500, "t_end": t + 501}) is None
    # the window gets it in pieces (Range), only with the key, only files it was given
    a_ = app.App()
    app.Handler.app = a_
    a_.media = {"k1": (fp, time.time())}
    srv = app.ThreadingHTTPServer(("127.0.0.1", 0), app.Handler)
    srv.daemon_threads = True
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    try:
        with httpx.Client(trust_env=False, timeout=10) as c:
            assert c.get(base + "/media?k=k1").status_code == 403
            c.get(base + "/?t=" + a_.token)
            full = c.get(base + "/media?k=k1")
            assert full.status_code == 200 and len(full.content) == os.path.getsize(fp)
            part = c.get(base + "/media?k=k1", headers={"Range": "bytes=10-19"})
            assert part.status_code == 206 and part.content == full.content[10:20] and part.headers["content-range"].endswith(f"/{len(full.content)}")
            assert c.get(base + "/media?k=k1", headers={"Range": "bytes=99999999-"}).status_code == 416
            assert c.get(base + "/media?k=other").status_code == 404
    finally:
        srv.shutdown()
        a_.closing.set()
    assert app.load_audio_index(os.path.join(str(tmp_path), "nothing.json")) == []
    with open(os.path.join(str(tmp_path), "bad.audio.json"), "w") as f:
        f.write('{"parts": [{"source": "them", "file": "../../x.ogg", "t0": 1, "dur": 2}]}')
    assert app.load_audio_index(os.path.join(str(tmp_path), "bad.audio.json")) == []      # no paths out of the folder


# ---- 6.14 review round: the cases the independent review found ---------------------------------
def test_subtitles_never_hang_on_a_long_word_and_escape_vtt(app):
    import time as _t
    link = "https://example.com/" + "x" * 400
    rows = [{"source": "them", "t0": 0.0, "t_end": 60.0, "text": link + " see <b> & --> here", "translation": "", "speaker": None}]
    t = _t.time()
    srt = app.srt_text(rows, 0.0, lambda r: "Them")
    assert _t.time() - t < 2 and "-->" not in "".join(c.split("\n", 2)[2] for c in srt.strip().split("\n\n"))
    for c in srt.strip().split("\n\n"):
        assert all(len(x) <= 2 * app.SUB_LINE for x in c.split("\n")[2:])

    class S:
        recording = None
        started = app.datetime.datetime.fromtimestamp(1_700_001_000.0)   # a real date: Windows cannot go back to 1970
        def ordered(self):
            return [{"source": "them", "t0": 1_700_001_001.0, "t_end": 1_700_001_003.0, "text": "a < b & c", "translation": "", "speaker": None}]
        def label(self, r):
            return "Them"
    vtt = app.export_bytes(S(), "vtt").decode()
    assert vtt.startswith("WEBVTT") and "a &lt; b &amp; c" in vtt and "00:00:01.000 --> 00:00:03.000" in vtt


def test_translation_is_cut_like_the_original(app):
    parts = app._split_like("one two three four five six seven eight nine ten", [30, 10])
    assert len(parts) == 2 and len(parts[0].split()) > len(parts[1].split()) and " ".join(parts).split() == \
        "one two three four five six seven eight nine ten".split()
    p2 = app._split_like("first part here, and the second part is here", [17, 26])
    assert p2[0].endswith(",")                                   # a comma near the cut is used
    assert app._split_like("a b", 3) == ["a b", "", ""] and app._split_like("", [1, 1]) == ["", ""]


def test_helper_works_while_a_link_is_downloading(app):
    a = app.App()
    try:
        job = app.LinkJob(a, "http://127.0.0.1:1/x")
        a.recording = job
        assert a.helper() is a.review                            # no crash: a link has no engine yet
    finally:
        a.recording = None
        a.closing.set()


def test_whisper_cpp_model_file_is_offered_and_counts_as_an_engine(app, tmp_path):
    import os
    base = str(tmp_path)
    srv = os.path.join(base, "whisper-server")
    open(srv, "w").write("x")
    d = os.path.join(base, "m")
    os.makedirs(d)
    open(os.path.join(d, "ggml-small-q5_1.bin"), "wb").write(b"lmgg" + b"\0" * 64)
    open(os.path.join(d, "notes.bin"), "wb").write(b"hello")
    old = os.environ.get("MA_WHISPER_SERVER")
    os.environ["MA_WHISPER_SERVER"] = srv
    a = app.App()
    try:
        r = a.api_browse(d, "model")
        assert [f["name"] for f in r["files"]] == ["ggml-small-q5_1.bin"]
        assert app.model_folder(os.path.join(d, "ggml-small-q5_1.bin"))[0].endswith("ggml-small-q5_1.bin")
        assert app.LOCAL.info()["engine"] is True                    # whisper.cpp alone is enough to run a model
    finally:
        a.closing.set()
        if old is None:
            os.environ.pop("MA_WHISPER_SERVER", None)
        else:
            os.environ["MA_WHISPER_SERVER"] = old


def test_meeting_sound_keeps_silence_while_paused_and_never_blocks(app, tmp_path):
    import os, time, wave
    np = app.np
    rec = app.MeetingAudio(os.path.join(str(tmp_path), "meeting_p.md"))
    rec.write("them", np.full(16000, 0.5, np.float32), 16000)
    rec.paused = True
    rec.write("them", np.full(16000, 0.5, np.float32), 16000)
    t = time.time()
    rec.write("them", np.full(16000, 0.5, np.float32), 16000)
    assert time.time() - t < 0.05                                  # the capture thread only queues
    rec.close()
    part = app.load_audio_index(os.path.join(str(tmp_path), "meeting_p.audio.json"))[0]
    fp = os.path.join(str(tmp_path), part["file"])
    if fp.endswith(".wav"):
        with wave.open(fp) as w:
            y = np.frombuffer(w.readframes(w.getnframes()), np.int16)
        assert np.abs(y[:16000]).mean() > 10000 and np.abs(y[16000:]).max() == 0


def test_ffmpeg_is_found_where_winget_puts_it_and_the_window_is_told(app, tmp_path):
    import os, shutil
    base = str(tmp_path)
    links = os.path.join(base, "Microsoft", "WinGet", "Links")
    os.makedirs(links)
    old_which, old_la = shutil.which, os.environ.get("LOCALAPPDATA")
    m4a = os.path.join(base, "talk.m4a")
    open(m4a, "wb").write(b"\0\0\0\x18ftypM4A " + b"\0" * 200)
    a = app.App()
    try:
        shutil.which = lambda name, *x, **k: None if name in ("ffmpeg", "ffprobe") else old_which(name, *x, **k)
        os.environ["LOCALAPPDATA"] = base
        assert app.ffmpeg_path() is None and app.needs_ffmpeg(m4a)
        a.missing_tasks = lambda *x: []
        r = a.api_recording(m4a)
        assert r["ok"] is False and r["error"] == "ffmpeg" and "m4a" in r["detail"] and a.recording is None
        assert "winget install Gyan.FFmpeg" in str(a.api_selfcheck())
        open(os.path.join(links, "ffmpeg.exe"), "wb").write(b"x")
        assert app.ffmpeg_path() == os.path.join(links, "ffmpeg.exe") and not app.needs_ffmpeg(m4a)
    finally:
        shutil.which = old_which
        if old_la is None:
            os.environ.pop("LOCALAPPDATA", None)
        else:
            os.environ["LOCALAPPDATA"] = old_la
        a.closing.set()


def test_link_problems_are_plain_and_stop_is_quick(app):
    import socket, threading, time
    assert app.link_problem("ERROR: [youtube] x: Private video. Sign in if you've been granted access").startswith("This video is private")
    assert "VPN" in app.link_problem("Unable to download webpage: <urlopen error [Errno 11001] getaddrinfo failed>")
    assert "robot" in app.link_problem("Sign in to confirm you’re not a bot")
    assert app.link_problem("something odd").startswith("The link could not be downloaded.")
    try:
        import yt_dlp  # noqa: F401
    except ImportError:
        return
    srv = socket.socket()
    srv.bind(("127.0.0.1", 0))
    srv.listen(5)                                               # accepts, never answers: a very slow site
    a = app.App()
    try:
        job = app.LinkJob(a, f"http://127.0.0.1:{srv.getsockname()[1]}/v.mp3")
        a.recording = job
        job.thread.start()
        time.sleep(0.8)
        t = time.time()
        a.api_recording(cancel=True)
        job.thread.join(5)
        assert not job.thread.is_alive() and time.time() - t < 3 and a.recording is None
        assert job.state["state"] == "cancelled"
    finally:
        srv.close()
        a.closing.set()


def test_play_button_knows_when_a_meeting_has_sound(app, tmp_path):
    import os, types
    np = app.np
    path = os.path.join(str(tmp_path), "meeting_h.md")
    s = types.SimpleNamespace(path=path, recording="")
    assert not app.session_has_audio(s) and not app.session_has_audio(None)
    rec = app.MeetingAudio(path)
    rec.write("them", np.full(16000, 0.1, np.float32), 16000)
    rec.close()
    assert app.session_has_audio(s)
    gone = types.SimpleNamespace(path=path, recording=os.path.join(str(tmp_path), "moved.mp3"))
    assert not app.session_has_audio(gone)


def test_vosk_model_is_downloaded_unpacked_and_found(app, tmp_path):
    import os, io, zipfile, threading, http.server, functools
    base = str(tmp_path)
    srv_dir = os.path.join(base, "site")
    os.makedirs(srv_dir)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:                         # the layout of a real Vosk model zip
        z.writestr("vosk-model-small-xx-0.1/am/final.mdl", b"m" * 5000)
        z.writestr("vosk-model-small-xx-0.1/conf/model.conf", b"--sample-frequency=16000\n")
        z.writestr("vosk-model-small-xx-0.1/graph/HCLr.fst", b"g" * 3000)
        z.writestr("../../evil.txt", b"never outside the models folder")
    open(os.path.join(srv_dir, "m.zip"), "wb").write(buf.getvalue())
    h = functools.partial(http.server.SimpleHTTPRequestHandler, directory=srv_dir)
    h.log_message = lambda *a: None
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), h)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    lib = os.path.join(base, "libvosk.so")
    open(lib, "wb").write(b"x")
    old = os.environ.get("MA_VOSK_LIB")
    os.environ["MA_VOSK_LIB"] = lib                              # the engine is "already there": only the model comes
    try:
        item = {"id": "vosk-model-small-xx-0.1", "url": f"http://127.0.0.1:{srv.server_address[1]}/m.zip", "mb": 1,
                "engine": "vosk", "ui": "stt", "lang": "xx", "note": ""}
        seen = []
        dl = app.ZipDownload(item, "", seen.append)
        dl.run()
        assert dl.state["state"] == "done", dl.state
        dest = os.path.join(app.MODELS_DIR, item["id"])
        assert app.is_vosk_dir(dest) and app.model_folder(dest)[0] == os.path.abspath(dest)
        assert app.model_folder(os.path.join(dest, "am", "final.mdl"))[0] == os.path.abspath(dest)   # a file inside works too
        assert not os.path.exists(os.path.join(app.MODELS_DIR, "evil.txt")) and not os.path.exists(os.path.join(base, "evil.txt"))
        assert not os.path.exists(dest + ".zip.part") and not os.path.exists(dest + ".tmp")
        assert app.model_label(dest).endswith("(Vosk)") and app.folder_mb(dest) == 0 and app.vosk_lang(dest) == "xx"
        m = [x for x in app.find_models() if x["path"] == os.path.abspath(dest)]
        assert m and m[0]["engine"] == "vosk" and m[0]["lang"] == "xx"
        assert any(x["id"].startswith("vosk-model-small-en") for x in app.vosk_catalog())
    finally:
        srv.shutdown()
        if old is None:
            os.environ.pop("MA_VOSK_LIB", None)
        else:
            os.environ["MA_VOSK_LIB"] = old


def test_vosk_real_engine_when_available(app):
    """Runs only where the Vosk engine and an English model are given (MA_VOSK_LIB, MA_VOSK_TEST_MODEL, MA_VOSK_TEST_WAV)."""
    import os, wave
    lib, model, wav = (os.environ.get(k, "") for k in ("MA_VOSK_LIB", "MA_VOSK_TEST_MODEL", "MA_VOSK_TEST_WAV"))
    if not (os.path.isfile(lib) and os.path.isdir(model) and os.path.isfile(wav)) or open(lib, "rb").read(4) == b"x":
        return
    np = app.np
    with wave.open(wav) as w:
        a = np.frombuffer(w.readframes(w.getnframes()), np.int16).astype(np.float32) / 32768
    m = app.VoskModel(model)
    try:
        r = m.transcribe(a, "en")
        assert "country" in r["text"].lower() and r["text"][0].isupper() and r["segments"][0]["end"] > 5
    finally:
        m.free()


def test_vosk_download_button_uses_the_zip_download(app, tmp_path):
    import os
    lib = os.path.join(str(tmp_path), "libvosk.so")
    open(lib, "wb").write(b"x")
    old_env, old_run, old_mrun = os.environ.get("MA_VOSK_LIB"), app.ZipDownload.run, app.ModelDownload.run
    os.environ["MA_VOSK_LIB"] = lib
    app.ZipDownload.run = app.ModelDownload.run = lambda self: None     # no real download here: only which path is taken
    a = app.App()
    try:
        r = a.api_local_download(id="vosk-model-small-fa-0.42")
        assert r["ok"] and isinstance(a.download, app.ZipDownload) and a.download.item["lang"] == "fa"
        a.download.thread.join(2)
        r = a.api_local_download(id="small")
        assert r["ok"] and isinstance(a.download, app.ModelDownload)
        a.download.cancel.set()
        a.download.thread.join(5)
        assert app._vosk_text("i think i ragazzi", "it") == "I think i ragazzi" and app._vosk_text("i am here", "en") == "I am here"
    finally:
        app.ZipDownload.run, app.ModelDownload.run = old_run, old_mrun
        if old_env is None:
            os.environ.pop("MA_VOSK_LIB", None)
        else:
            os.environ["MA_VOSK_LIB"] = old_env
        a.closing.set()


def test_nvidia_support_is_downloaded_checked_and_found(app, tmp_path):
    import os, io, json, zipfile, hashlib, threading, http.server, shutil
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:                         # the layout of the real wheel
        z.writestr("nvidia/cublas/bin/cublas64_12.dll", b"A" * 4000)
        z.writestr("nvidia/cublas/bin/cublasLt64_12.dll", b"B" * 9000)
        z.writestr("nvidia/cublas/bin/nvblas64_12.dll", b"C" * 100)
        z.writestr("nvidia/cublas/lib/x64/cublas.lib", b"D" * 100)
    wheel = buf.getvalue()
    state = {"sha": hashlib.sha256(wheel).hexdigest(), "hits": 0}

    class H(http.server.BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_GET(self):
            if self.path.endswith("/json"):
                body = json.dumps({"urls": [{"filename": "nvidia_cublas_cu12-12.9.1.4-py3-none-win_amd64.whl",
                                             "url": f"http://127.0.0.1:{self.server.server_address[1]}/w.whl",
                                             "size": len(wheel), "digests": {"sha256": state["sha"]}}]}).encode()
                self.send_response(200); self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)
                return
            state["hits"] += 1
            start = int(self.headers.get("Range", "bytes=0-")[6:].split("-")[0] or 0)
            part = wheel[start:]
            if start and state["hits"] == 1:                     # never asked to continue the first time
                raise AssertionError("unexpected range")
            self.send_response(206 if start else 200)
            if start:
                self.send_header("Content-Range", f"bytes {start}-{len(wheel) - 1}/{len(wheel)}")
            cut = len(part) // 2 if state["hits"] == 1 else len(part)   # the first time the connection breaks halfway
            self.send_header("Content-Length", str(len(part)))
            self.end_headers()
            self.wfile.write(part[:cut])

    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    old = os.environ.get("MA_PYPI")
    os.environ["MA_PYPI"] = f"http://127.0.0.1:{srv.server_address[1]}/pypi"
    try:
        shutil.rmtree(app.gpu_pack_dir(), ignore_errors=True)
        assert not app.gpu_pack_ready()
        dl = app.GpuPackDownload("", lambda st: None)
        dl.run()
        assert dl.state["state"] == "done", dl.state
        assert app.gpu_pack_ready() and state["hits"] >= 2              # it continued after the break
        assert sorted(os.listdir(app.gpu_pack_dir())) == ["cublas64_12.dll", "cublasLt64_12.dll"]   # only the two it needs
        assert not os.path.exists(dl.part)
        assert app.gpu_pack_dir() in os.environ["PATH"].split(os.pathsep)
        shutil.rmtree(app.gpu_pack_dir())
        state["sha"], state["hits"] = "0" * 64, 5                        # a damaged file is refused and thrown away
        dl = app.GpuPackDownload("", lambda st: None)
        dl.run()
        assert dl.state["state"] == "error" and "SHA-256" in dl.state["error"] and not app.gpu_pack_ready()
        assert not os.path.exists(dl.part)
    finally:
        srv.shutdown()
        if old is None:
            os.environ.pop("MA_PYPI", None)
        else:
            os.environ["MA_PYPI"] = old


def test_nvidia_support_remove_and_notes(app):
    import os, shutil
    d = app.gpu_pack_dir()
    os.makedirs(d, exist_ok=True)
    for f in app.GPU_PACK["dlls"]:
        open(os.path.join(d, f), "wb").write(b"x")
    assert app.gpu_pack_ready()
    a = app.App()
    try:
        r = a.api_gpu_pack(remove=True)
        assert r["ok"] and not app.gpu_pack_ready() and not os.path.exists(d)
        os.makedirs(d)
        for f in app.GPU_PACK["dlls"]:
            open(os.path.join(d, f), "wb").write(b"x")
        open(os.path.join(d, "remove.flag"), "w").close()       # removed at the next start when it was in use
        app._GPU.pop("cleaned", None)
        assert not app.gpu_pack_ready() and not os.path.exists(d)
    finally:
        shutil.rmtree(d, ignore_errors=True)
        app._GPU.pop("cleaned", None)
        a.closing.set()


def test_audit_617_small_fixes(app, tmp_path):
    import os, time, types
    np = app.np
    # a mixed-slash network path is refused too
    assert app.is_network_path("\\/server/share/a.mp3") and app.is_network_path("/\\server\\x")
    assert not app.is_network_path("C:\\Users\\a.mp3")
    a = app.App()
    try:
        # the proxy password never goes to the window, and sending the masked text back keeps the real one
        a.api_save_config(proxy="http://user:secret@127.0.0.1:8080")
        shown = a.public_config()["proxy"]
        assert "secret" not in shown and "***" in shown
        a.api_save_config(proxy=shown)
        assert a.cfg["proxy"] == "http://user:secret@127.0.0.1:8080"
        a.api_save_config(proxy="")
        # a local model that is gone is not 'ready'
        a.cfg.update(stt_provider="local", local_model=os.path.join(str(tmp_path), "gone"), api_key="")
        st = a.task_status()["stt"]
        assert not st["ok"] and "cannot be used" in st["why"]
        # an earlier meeting replaced by a new one is still saved when late translations arrive
        s1 = app.Session(app.sanitize({}))
        e = s1.add("them", time.time(), time.time() + 1, "Last question?", "en", [1])
        s1.save_if_dirty()
        a._keep_saving(s1)
        s1.update(e["id"], translation="سؤال آخر؟", tr_state="done")
        for s in list(a.older):
            s.save_if_dirty()
        assert "سؤال آخر؟" in open(s1.path, encoding="utf-8").read()
    finally:
        a.closing.set()
    # the meeting sound: a second recorder (after a crash) never writes over the first files, and the index is kept current
    path = os.path.join(str(tmp_path), "meeting_c.md")
    r1 = app.MeetingAudio(path)
    r1._indexed = 0
    r1.write("them", np.full(16000, 0.3, np.float32), 16000)
    time.sleep(1.0)
    assert app.load_audio_index(os.path.join(str(tmp_path), "meeting_c.audio.json"))      # written while recording
    first = [f for f in os.listdir(str(tmp_path)) if f.startswith("meeting_c.them")]
    r1.files = None                                                   # 'crash': never closed
    r2 = app.MeetingAudio(path)
    r2.parts = []                                                     # as if the index were lost
    r2.write("them", np.full(16000, 0.3, np.float32), 16000)
    r2.close()
    names = sorted(f for f in os.listdir(str(tmp_path)) if f.startswith("meeting_c.them"))
    assert len(names) == 2 and first[0] in names                      # the old file is still there, untouched


def test_audit_616_settings_reset_links_and_dropped_engine(app):
    import os, json
    # a v3 settings file: only the kept keys come over; odd values are fixed
    with open(app.CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump({"config_version": 3, "api_key": "gsk_x", "local_model": "C:/old/model", "save_audio": "yes"}, f)
    c = app.load_config()
    assert c["api_key"] == "gsk_x" and c["local_model"] == "" and c["save_audio"] is False
    c2 = app.sanitize({"save_audio": "yes", "local_device": "gpu"})
    assert c2["save_audio"] in (True, False) and c2["local_device"] == "auto"
    # config.json missing: the backup copy is used, and the note says "missing", not "damaged"
    with open(app.CONFIG_PATH + ".bak", "w", encoding="utf-8") as f:
        json.dump({"config_version": app.CONFIG_VERSION, "api_key": "gsk_old"}, f)
    os.remove(app.CONFIG_PATH)
    app.CONFIG_NOTE.clear()
    c3 = app.load_config()                                      # missing (OneDrive, antivirus): the backup keeps the keys
    assert c3["api_key"] == "gsk_old" and any("missing" in n for n in app.CONFIG_NOTE)
    os.remove(app.CONFIG_PATH + ".bak")
    # link messages: a busy site is not "removed"; blocked sites get the VPN advice
    assert "busy" in app.link_problem("HTTP Error 503: Service Unavailable")
    assert "busy" in app.link_problem("HTTP Error 429: Too Many Requests")
    assert "VPN" in app.net_hint("model.bin: cannot reach the site", "huggingface.co")
    assert app.net_hint("model.bin is missing on the site", "x") == "model.bin is missing on the site"
    # a whisper.cpp engine that was dropped is never started again by a request still running
    srv = app.CppServer("x.bin", 1)
    srv.close()
    try:
        srv._restart()
        assert False
    except RuntimeError as e:
        assert "closed" in str(e)
    # Stop on a recording restarts whisper.cpp later, but is never counted as a crash
    srv2 = app.CppServer("x.bin", 1)
    srv2.start = lambda *a, **k: True
    now = app.time.time()
    srv2.restarts = [now - 5, now - 4, now - 3]
    srv2.interrupt()
    srv2._restart()                                              # a 4th crash would raise; a planned one does not
    assert len(srv2.restarts) == 3
    # deleting never touches a network path
    a = app.App()
    try:
        assert a.api_local_delete(path=r"\\server\share\model")["ok"] is False
    finally:
        a.closing.set()


def test_features_can_be_turned_off(app):
    c = app.sanitize({"features": {"answers": False, "coach": "false", "bogus": False, "translate": True}})
    assert c["features"] == {"answers": False, "coach": False, "translate": True}
    assert not app.feat(c, "answers") and app.feat(c, "screen") and not app.coach_on(dict(c, coach=True))
    a = app.App()
    try:
        a.cfg["features"] = {"translate": False, "answers": False, "coach": False}
        assert a.needed_tasks() == ("stt", "ans")                   # prep, summary and screen reading still use answers
        a.cfg["features"] = {"translate": False, "answers": False, "coach": False, "prep": False, "after": False, "screen": False}
        assert a.needed_tasks() == ("stt", "ans")                   # 6.19: say mode and full notes use answers too
        a.cfg["features"].update(say=False, notes=False)
        assert a.needed_tasks() == ("stt",)                         # only speech to text needs a service
        assert a.needed_tasks(for_file=True) == ("stt",)
        a.cfg["features"] = {}
        assert a.needed_tasks() == ("stt", "tr", "ans")
        a.cfg["features"] = {"answers": False, "screen": False, "after": False, "prep": False, "files": False,
                             "overlay": False, "coach": False, "say": False, "notes": False, "marks": False}
        for name, kw in (("answer_now", {}), ("screen", {}), ("summary", {}), ("feedback", {}), ("job_prep", {"ad": "x" * 60}),
                         ("practice_questions", {}), ("say", {"text": "hello"}), ("recording", {"path": "x.mp3"}),
                         ("recording_link", {"url": "https://example.com/v"}), ("overlay", {"on": True}),
                         ("say_mode", {"on": True}), ("mark", {}), ("notes", {"text": "x"}), ("notes_full", {}),
                         ("glossary_suggest", {})):
            r = a.call(name, kw)
            assert r["ok"] is False and r.get("feature") and "Setup › Features" in r["error"], (name, r)
        assert a.call("recording", {"cancel": True})["ok"] is True          # stopping always works
        pub = a.public_config()
        assert pub["tasks"]["ans"]["used"] is False and pub["tasks"]["tr"]["used"] is True
    finally:
        a.closing.set()


def test_lines_follow_the_feature_switches(app):
    import time
    eng, sess, calls = _engine(app)
    tr = []
    eng._submit_tr = lambda eid: tr.append(eid) or True
    try:
        eng.cfg["features"] = {"translate": False, "answers": False}
        _say(app, eng, "How would you design a backup strategy for a very large database?")
        time.sleep(1.5)
        e = sess.ordered()[-1]
        assert e["tr_state"] == "off" and e["ans_state"] == "off" and not tr and not calls   # nothing sent
        eng.cfg["features"] = {}
        _say(app, eng, "And how would you test that the backups really work?")
        time.sleep(1.8)
        e = sess.ordered()[-1]
        assert e["tr_state"] != "off" and tr and calls                                     # on again: both run
    finally:
        eng.stop_event.set()


# ---- 6.19: say it in my language, marks, speaker names, my notes, glossary, meeting reminder --------------
class _RecHub:
    def __init__(self):
        self.events, self.toasts = [], []
    def publish(self, kind, **k): self.events.append((kind, k))
    def toast(self, level, text): self.toasts.append((level, text))


def test_marks_names_and_notes_reach_the_files(app):
    cfg = app.sanitize({})
    s = app.Session(cfg)
    t = time.time()
    a = s.add("them", t, t + 2, "What is your notice period?", "en", [1], speaker=1)
    b = s.add("them", t + 3, t + 5, "And your salary expectation?", "en", [2], speaker=2)
    s.add("me", t + 6, t + 8, "Three months.", "en", [3])
    assert s.mark(None)["id"] == s.ordered()[-1]["id"]            # F7: the newest line, always "on"
    assert s.mark(None)["marked"] is True                         # pressing again never removes it
    assert s.mark(a["id"])["marked"] is True and s.mark(a["id"])["marked"] is False and s.mark(a["id"])["marked"] is True
    assert s.mark(99999) is None
    assert s.set_name("them:1", "  John   Smith ") and s.label(s.get(a["id"])) == "John Smith"
    assert s.label(s.get(b["id"])) != "John Smith"                 # only that speaker
    assert not s.set_name("me", "X") and not s.set_name("bad key", "X") and not s.set_name("them:1234", "X")
    assert not s.set_name("them:1\n", "X")
    assert s.set_name("them:2", "")                                # empty = back to the default name
    assert app.speaker_key({"source": "them", "who": "Person 2"}) == "who:Person 2"
    assert app.speaker_key({"source": "them", "speaker": None}) == "them:0"
    s.set_notes(mine="notice period?\n- salary")
    s.save_if_dirty()
    md = open(s.path, encoding="utf-8").read()
    assert "★ John Smith:** What is your notice period?" in md and "## Marked moments" in md and "## My notes" in md
    txt = app.export_bytes(s, "txt").decode("utf-8")
    assert "== Marked moments ==" in txt and "★ [" in txt and "== My notes ==" in txt and "John Smith" in txt
    with zipfile.ZipFile(io.BytesIO(app.export_bytes(s, "docx"))) as z:
        doc = z.read("word/document.xml").decode("utf-8")
    assert "Marked moments" in doc and "John Smith" in doc and "salary" in doc
    back = app.Session.load_last(cfg)
    assert back.names == {"them:1": "John Smith"} and back.my_notes.startswith("notice period?")
    assert [r["marked"] for r in back.ordered()] == [True, False, True]
    # damaged or hostile values from the file are cleaned
    assert app.clean_names({"them:1": "A", "x": "B", "who:" + "z" * 80: "C", "them:2": 5, "them:3": " " * 3}) == {"them:1": "A"}


def test_mark_rename_notes_actions(app):
    a = app.App()
    try:
        a.hub = _RecHub()
        assert a.api_mark()["ok"] is False                        # no transcript yet
        eng, sess, _ = _engine(app)
        a.review = eng
        t = time.time()
        e = sess.add("them", t, t + 2, "Tell me about yourself.", "en", [1])
        r = a.api_mark()
        assert r["ok"] and r["id"] == e["id"] and r["marked"] and a.hub.toasts and "Marked" in a.hub.toasts[-1][1]
        assert a.api_mark(entry_id=str(e["id"]))["marked"] is False      # the flag on a line switches it
        assert a.api_mark(entry_id="abc")["ok"] is False
        assert a.api_mark(entry_id=True)["ok"] is False and a.api_mark(entry_id=3.9)["ok"] is False
        assert a.api_mark(entry_id=e["id"], on="false")["marked"] is True          # only a real true/false is taken
        assert a.api_mark(entry_id=e["id"], on=False)["marked"] is False
        r = a.api_rename_speaker(key="them:0", name="Interviewer")
        assert r["ok"] and r["names"] == {"them:0": "Interviewer"} and sess.label(sess.get(e["id"])) == "Interviewer"
        assert ("names", {"names": {"them:0": "Interviewer"}, "session_file": sess.path}) in a.hub.events
        assert a.api_rename_speaker(key="me", name="X")["ok"] is False
        assert a.api_notes(text="n" * 30000)["ok"] and len(sess.my_notes) == 20000
        r = a.api_notes(text="other", session_file="C:/elsewhere/meeting.md")
        assert r["ok"] is False and r["stale"] and sess.my_notes == "n" * 20000          # never into the wrong meeting
        assert a.api_notes(text="kept", session_file=sess.path)["ok"] and sess.my_notes == "kept"
        assert a.api_rename_speaker(key="them:0", name="X", session_file="other.md")["ok"] is False
        assert os.path.isfile(sess.path)                          # saved at once (no meeting is running)
        a.cfg["features"] = {"marks": False, "notes": False}
        assert a.api_mark()["feature"] == "marks" and a.api_notes(text="x")["feature"] == "notes"
        assert a.api_rename_speaker(key="them:0", name="")["ok"] is True      # names are part of the transcript itself
        eng.stop_event.set()
    finally:
        a.closing.set()


def test_say_mode_routes_my_speech_to_a_sentence(app):
    import numpy as np
    eng, sess, calls = _engine(app)
    hub = _RecHub()
    eng.hub = hub
    eng.app.hub = hub
    eng.cfg["my_language"] = "fa"
    eng.cfg["languages"] = ["en"]
    try:
        seen = {}

        class FakeSTT:
            def transcribe(self, audio, model, lang, prompt):
                seen["lang"] = lang
                return {"text": "بگو که من سه سال تجربه دارم", "language": "persian", "segments": []}
        eng.app.get_api = lambda pid: FakeSTT()
        eng.app.note_request = lambda: None
        eng.say = lambda text, lang="", tone="natural": {"text": "I have three years of experience.", "meaning": "سه سال تجربه دارم", "lang": "en"}
        eng._submit = lambda fn, *a: fn(*a) or True
        eng.say_mode = True
        audio = np.zeros(16000, dtype=np.float32)
        eng.on_audio("segment", "me", 77, audio=audio, t0=time.time() - 2, t_end=time.time(), dur=1.0)
        with eng.cv:
            job = eng.pending.pop()
        assert job["say"] is True and job["source"] == "me"
        eng._process(job, ("groq", "whisper-large-v3"))
        assert seen["lang"] == "fa"                               # recognised in MY language, not the meeting language
        assert sess.ordered() == []                               # never written into the transcript
        done = [k for kind, k in hub.events if kind == "say" and k.get("state") == "done"]
        assert done and done[-1]["text"] == "I have three years of experience." and done[-1]["heard"].startswith("بگو")
        # the other side is not affected, and with say mode off my words are a line again
        eng.on_audio("segment", "them", 78, audio=audio, t0=time.time() - 2, t_end=time.time(), dur=1.0)
        with eng.cv:
            assert not eng.pending[-1].get("say")
            eng.pending.clear()
        eng.say_mode = False
        eng.on_audio("segment", "me", 79, audio=audio, t0=time.time() - 2, t_end=time.time(), dur=1.0)
        with eng.cv:
            assert not eng.pending[-1].get("say")
            eng.pending.clear()
        # a live service never writes my line while say mode is on
        eng.say_mode = True
        eng.live_final("me", 80, time.time() - 1, time.time(), "garbled words")
        assert sess.ordered() == []
        # nothing understood: a clear message, nothing sent
        eng.say_speech("  ")
        assert hub.events[-1][1]["state"] == "error"
        # an older sentence that finishes after a newer one is not shown over it
        hub.events.clear()
        eng._say_speech("new", 9)
        eng._say_speech("old", 8)
        assert [k["heard"] for kind, k in hub.events if kind == "say"] == ["new"]
        # a say piece that cannot be written down tells the card
        eng._drop({"segs": [81], "source": "me", "say": True})
        assert hub.events[-1][0] == "say" and hub.events[-1][1]["state"] == "error"
    finally:
        eng.stop_event.set()


def test_say_mode_action_checks(app):
    a = app.App()
    try:
        assert a.api_say_mode(on=True) == {"ok": False, "error": "Start the meeting first."}
        assert a.api_say_mode(on=False)["ok"] is True                    # turning off always works
        a.cfg["features"] = {"say": False}
        assert a.api_say_mode(on=True)["feature"] == "say"
        eng, sess, _ = _engine(app)
        a.engine = eng
        a.cfg["features"] = {}
        a.captures = []
        assert "microphone" in a.api_say_mode(on=True)["error"].lower()
        class Cap: source = "me"
        a.captures = [Cap()]
        a.hub = _RecHub()
        assert a.api_say_mode(on=True) == {"ok": True, "on": True} and eng.say_mode
        assert any("mute" in t.lower() for _, t in a.hub.toasts)
        assert a.api_say_mode() == {"ok": True, "on": False} and not eng.say_mode     # the key switches it
        eng.say_mode = True
        a.cfg["features"] = {"say": False}
        a._apply_features()
        assert not eng.say_mode                                           # turned off in Features: stops at once
        eng.stop_event.set()
    finally:
        a.closing.set()


def test_glossary_merge_and_suggest(app):
    text, added = app.merge_glossary("OSPF, BGP\nkubectl", ["bgp", "Data Guard", "  RMAN ", "x" * 50, "Data guard", ""])
    assert added == ["Data Guard", "RMAN"] and text == "OSPF, BGP\nkubectl\nData Guard\nRMAN"
    assert app.merge_glossary("", ["k8s"]) == ("k8s", ["k8s"])
    assert app.merge_glossary("A", []) == ("A", [])
    full = "\n".join(f"t{i}" for i in range(app.GLOSSARY_MAX_TERMS))
    assert app.merge_glossary(full, ["new"])[1] == []                    # the limit is kept
    eng, sess, _ = _engine(app)
    try:
        eng.cfg["job_ad"] = "We need Oracle RMAN and Data Guard, Kubernetes, Terraform."
        eng.run_chat = lambda *a, **k: 'Here: ["RMAN", "Data Guard", "rman", "Kubernetes", "' + "y" * 60 + '"]'
        assert eng.glossary_suggest() == ["RMAN", "Data Guard", "Kubernetes"]
        eng.run_chat = lambda *a, **k: "RMAN\n- Terraform\n"
        assert eng.glossary_suggest() == ["RMAN", "Terraform"]
    finally:
        eng.stop_event.set()
    a = app.App()
    try:
        a.cfg["job_ad"] = ""
        a.cfg["about_me"] = ""
        assert "job advert" in a.api_glossary_suggest()["error"]
    finally:
        a.closing.set()


def test_meeting_windows_are_recognised(app):
    f = app.meeting_apps
    assert f(["Zoom Meeting", "Inbox - Outlook"]) == {"Zoom"}
    assert f(["Meeting with Sara | Microsoft Teams", "Chat | Microsoft Teams"]) == {"Teams"}
    assert f(["Meet - abc-defg-hij - Google Chrome"]) == {"Google Meet"}
    assert f(["Chat | Microsoft Teams", "Zoom Workplace", "Google Docs", "meet the team - Notepad"]) == set()
    assert f(["Cisco Webex Meetings"]) == {"Webex"}
    assert f(["Chat | Weekly Meeting | Microsoft Teams", "Webex Meetings pricing - Google Chrome",     # pages and chats
              "Zoom Meeting tips - YouTube - Google Chrome", "Weekly Meeting notes - Microsoft\u200b Edge"]) == set()
    assert f(["Meet - abc-defg-hij and 2 more pages - Personal - Microsoft\u200b Edge"]) == {"Google Meet"}
    old = os.environ.get("MA_FAKE_WINDOWS")
    os.environ["MA_FAKE_WINDOWS"] = "Zoom Meeting;;Notes"
    try:
        assert app.window_titles() == ["Zoom Meeting", "Notes"]
    finally:
        if old is None:
            os.environ.pop("MA_FAKE_WINDOWS")
        else:
            os.environ["MA_FAKE_WINDOWS"] = old


def test_meeting_watch_offers_start_once(app, tmp_path):
    import threading
    p = os.path.join(tmp_path, "w.txt")
    open(p, "w").write("Inbox\n")
    old = {k: os.environ.get(k) for k in ("MA_FAKE_WINDOWS", "MA_WATCH_SECS")}
    os.environ["MA_FAKE_WINDOWS"], os.environ["MA_WATCH_SECS"] = p, "0.05"
    a = app.App()
    try:
        a.hub = _RecHub()
        th = threading.Thread(target=a._meeting_watch, daemon=True)
        th.start()
        time.sleep(0.2)
        open(p, "w").write("Inbox\nZoom Meeting\n")
        time.sleep(0.3)
        found = [k for kind, k in a.hub.events if kind == "meeting_found"]
        assert found == [{"app": "Zoom"}]
        open(p, "w").write("Inbox\n")
        time.sleep(0.2)
        open(p, "w").write("Zoom Meeting\n")                       # opened again soon after: not offered again
        time.sleep(0.3)
        assert len([1 for kind, _ in a.hub.events if kind == "meeting_found"]) == 1
        a.cfg["features"] = {"detect": False}
        open(p, "w").write("Meeting with Sara | Microsoft Teams\n")
        time.sleep(0.3)
        assert len([1 for kind, _ in a.hub.events if kind == "meeting_found"]) == 1     # turned off
        a.cfg["features"] = {}
        a.running = True
        time.sleep(0.3)
        assert len([1 for kind, _ in a.hub.events if kind == "meeting_found"]) == 1     # a meeting is already running
    finally:
        a.closing.set()
        for k, v in old.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def test_summary_asks_for_decisions_and_actions_and_uses_marks(app):
    eng, sess, _ = _engine(app)
    hub = _RecHub()
    eng.hub = hub
    try:
        t = time.time()
        sess.add("them", t, t + 2, "We decided to move the release to Friday.", "en", [1])
        e = sess.add("me", t + 3, t + 5, "I will send the test report by Thursday.", "en", [2])
        sess.mark(e["id"])
        sess.set_notes(mine="release friday")
        got = {}

        def chat(chain, msgs, *a, **k):
            got["system"], got["user"] = msgs[0]["content"], msgs[1]["content"]
            return "## Summary\nok"
        eng.chat_patient = chat
        eng.cfg["my_language"] = "en"
        eng._summary(sess.ordered())
        assert "## Decisions" in got["system"] and "## Action items" in got["system"] and "by when" in got["system"]
        assert "★ " in got["user"] and "release friday" in got["user"]
        eng.cfg["my_language"] = "fa"
        eng._summary(sess.ordered())
        assert "## تصمیم‌ها" in got["system"]
        # full notes: my notes are the outline
        eng._notes_full(sess.ordered())
        assert "My notes:\nrelease friday" in got["user"] and sess.notes_full == "## Summary\nok"
        assert hub.events[-1] == ("notes_full", {"session_file": sess.path, "state": "done", "text": "## Summary\nok"})
    finally:
        eng.stop_event.set()


def test_filler_words_and_pace(app):
    rows = [{"source": "me", "t0": 0, "t_end": 60, "text": "So, like, I actually like Python. Um, you know, I basically, like, automated it. " + "word " * 120}]
    st = app.talk_stats(rows)
    assert st["filler_list"]["like"] == 2                            # "I like Python" is not a filler
    assert list(st["filler_list"])[0] == "like" and st["fillers"] == 6
    assert st["fillers_per_min"] == 6.0 and st["pace"] == "good"
    fast = app.talk_stats([{"source": "me", "t0": 0, "t_end": 30, "text": "word " * 100}])
    assert fast["pace"] == "fast" and app.talk_stats([{"source": "me", "t0": 0, "t_end": 5, "text": "hi"}])["pace"] is None


def test_new_parts_are_off_for_someone_without_the_answers_service(app):
    import json as _j
    with open(app.CONFIG_PATH, "w", encoding="utf-8") as f:
        _j.dump({"config_version": 5, "api_key": "gsk_x", "features": {"answers": False, "after": False, "translate": False}}, f)
    c = app.load_config()
    assert c["features"]["say"] is False and c["features"]["notes"] is False and app.feat(c, "marks") and app.feat(c, "detect")
    with open(app.CONFIG_PATH, "w", encoding="utf-8") as f:
        _j.dump({"config_version": 5, "api_key": "gsk_x", "features": {"translate": False}}, f)
    c = app.load_config()
    assert app.feat(c, "say") and app.feat(c, "notes")
    with open(app.CONFIG_PATH, "w", encoding="utf-8") as f:              # a 6.19 file is taken as it is
        _j.dump({"config_version": 6, "api_key": "gsk_x", "features": {"answers": False, "after": False}}, f)
    assert app.feat(app.load_config(), "say")


# ---- 6.19: "is it really hidden from screen sharing?" (version detection + self-test judge) ----
def test_hide_capture_support_and_build(app):
    import sys
    assert app.WIN_HIDE_BUILD == 19041
    b = app.windows_build()
    if sys.platform != "win32":
        assert b == 0 and app.capture_hide_supported() is False                    # not Windows: no build, no hiding
    else:
        assert b > 0 and app.capture_hide_supported() is (b >= app.WIN_HIDE_BUILD)  # the real build decides
    assert app.windows_build() == b                                                 # cached, stable
    pub = app.App()
    try:
        c = pub.public_config()
        assert c["win_build"] == b and c["hide_capture_supported"] == (b >= app.WIN_HIDE_BUILD)
        assert isinstance(c["hide_possible"], bool)
    finally:
        pub.closing.set()


def test_judge_hide(app):
    bg = [(210, 210, 210)] * 49       # the background behind the window (bright)
    win = [(20, 110, 240)] * 49       # the window's own look
    black = [(0, 0, 0)] * 49
    assert app.judge_hide(win, win)[0] == "visible"        # the share shows the window
    assert app.judge_hide(bg, win)[0] == "hidden"          # the share shows the background instead
    assert app.judge_hide(black, win)[0] == "black"        # the share is a black box, the window is not
    assert app.judge_hide(bg, black)[0] == "hidden"        # a dark window is not mistaken for the black-box mode
    assert app.judge_hide(black, black)[0] != "black"      # (both dark: cannot claim the black-box mode)
    assert app.judge_hide([], [])[0] == "unknown"
    r, match, blk = app.judge_hide(bg, win)
    assert 0.0 <= match <= 1.0 and blk == 0.0


def test_hide_test_action_guards(app):
    a = app.App()
    try:
        r = a.api_hide_test()
        assert r["ok"] is False and ("Windows" in r["error"] or "hidden window" in r["error"])
    finally:
        a.closing.set()
